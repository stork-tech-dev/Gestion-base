"""Lectura y generación de planillas Excel (.xlsx) con openpyxl.

Única vía para importar/exportar tablas simples: el dólar por proveedor
(módulo 03), la importación inicial (módulo 08) y los reportes (módulo 09)."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.errores import ReglaDeNegocio
from app.core.utils import normalizar_texto, sin_tildes

TAMANO_MAXIMO_BYTES = 5 * 1024 * 1024
FILAS_MAXIMAS = 5000
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Un .xlsx es un zip: 5 MB comprimidos pueden ser gigas descomprimidos (zip
# bomb). Se suma el tamaño declarado de cada entrada antes de abrirlo; zipfile
# nunca entrega más bytes que los declarados, así que el límite es efectivo.
TAMANO_DESCOMPRIMIDO_MAXIMO = 50 * 1024 * 1024
# Un texto que empiece con alguno de estos caracteres Excel lo puede tomar como
# fórmula (inyección de fórmulas / CSV injection).
PREFIJOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def clave_columna(texto) -> str:
    """Encabezado normalizado: sin tildes, en minúscula y con espacios colapsados."""
    if texto is None:
        return ""
    return sin_tildes(normalizar_texto(str(texto)) or "").lower()


def _fila_vacia(valores) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in valores)


def leer_tabla(contenido: bytes, requeridas, opcionales=()) -> list[tuple[int, dict]]:
    """Filas de la primera hoja como (número de fila en Excel, {columna: valor}).

    Las columnas se ubican por nombre de encabezado (cualquier orden, sin
    importar mayúsculas ni tildes; las columnas extra se ignoran). `requeridas`
    y `opcionales` se pasan ya normalizadas (ej. "proveedor_id"). Las fórmulas
    sin valor calculado (archivos no guardados por Excel) se leen como vacías."""
    if not contenido:
        raise ReglaDeNegocio("El archivo está vacío.")
    if len(contenido) > TAMANO_MAXIMO_BYTES:
        raise ReglaDeNegocio("El archivo supera los 5 MB.")
    try:
        with ZipFile(BytesIO(contenido)) as comprimido:
            descomprimido = sum(entrada.file_size for entrada in comprimido.infolist())
    except Exception as exc:
        # Bytes no confiables: además de BadZipFile, un zip corrupto puede
        # levantar NotImplementedError (versión no soportada) o
        # UnicodeDecodeError (nombre de entrada corrupto), entre otros.
        raise ReglaDeNegocio("El archivo no es un Excel (.xlsx) válido.") from exc
    if descomprimido > TAMANO_DESCOMPRIMIDO_MAXIMO:
        raise ReglaDeNegocio("El archivo es demasiado grande una vez descomprimido.")
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:
        # Esta función parsea bytes no confiables; cualquier falla de parseo es "archivo inválido".
        raise ReglaDeNegocio("El archivo no es un Excel (.xlsx) válido.") from exc
    try:
        filas = libro.worksheets[0].iter_rows(values_only=True)
        encabezado = next(filas, None)
        if encabezado is None:
            raise ReglaDeNegocio("El archivo no tiene encabezados.")
        posiciones: dict[str, int] = {}
        for indice, celda in enumerate(encabezado):
            clave = clave_columna(celda)
            if not clave:
                continue
            if clave not in posiciones:
                posiciones[clave] = indice
            # "Dólar Actual" (encabezado a mano) y "dolar_actual" (nombre de
            # columna interno, el que también usa la plantilla generada) son
            # la misma columna: se indexa también con "_" para que ambas
            # formas encuentren la posición.
            con_guion = clave.replace(" ", "_")
            if con_guion not in posiciones:
                posiciones[con_guion] = indice
        faltan = [c for c in requeridas if c not in posiciones]
        if faltan:
            raise ReglaDeNegocio(f"Faltan columnas: {', '.join(faltan)}")
        resultado = []
        for numero, valores in enumerate(filas, start=2):
            if valores is None or _fila_vacia(valores):
                continue
            if len(resultado) >= FILAS_MAXIMAS:
                raise ReglaDeNegocio(f"El archivo supera las {FILAS_MAXIMAS} filas.")
            fila = {}
            for columna in (*requeridas, *opcionales):
                indice = posiciones.get(columna)
                fila[columna] = valores[indice] if indice is not None and indice < len(valores) else None
            resultado.append((numero, fila))
        return resultado
    except ReglaDeNegocio:
        raise
    except Exception as exc:
        raise ReglaDeNegocio("El archivo no es un Excel (.xlsx) válido.") from exc
    finally:
        libro.close()


def _textos_como_texto(hoja) -> None:
    """Fuerza como texto (tipo "s") toda celda str que Excel podría ejecutar
    como fórmula: openpyxl marca como fórmula cualquier str que empiece con
    "=". Así un nombre como '=HYPERLINK(...)' se guarda y se lee tal cual."""
    for fila in hoja.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.startswith(PREFIJOS_FORMULA):
                celda.data_type = "s"


def generar_tabla(hoja: str, columnas: list[str], filas) -> bytes:
    """Libro de una sola hoja (encabezados en negrita, anchos ajustados a los
    datos, panel congelado bajo el encabezado y protegida contra fórmulas con
    `_textos_como_texto`). Única definición: la reusan `generar_plantilla` y
    cualquier reporte que exporte una tabla simple a Excel (ej. auditorías de
    inventario)."""
    libro = Workbook()
    datos = libro.active
    datos.title = hoja
    datos.append(list(columnas))
    for fila in filas:
        datos.append(list(fila))
    for celda in datos[1]:
        celda.font = Font(bold=True)
    for indice, columna in enumerate(columnas, start=1):
        largo = max([len(str(columna))] + [len(str(f[indice - 1] or "")) for f in filas if len(f) >= indice])
        datos.column_dimensions[get_column_letter(indice)].width = min(max(12, largo + 4), 60)
    datos.freeze_panes = "A2"
    _textos_como_texto(datos)

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def generar_plantilla(hoja: str, columnas: list[str], instrucciones: list[list], filas=None) -> bytes:
    """Libro con la hoja de datos (`generar_tabla`, con filas precargadas
    opcionales) y una hoja "Instrucciones" (la primera fila en negrita)."""
    libro = load_workbook(BytesIO(generar_tabla(hoja, columnas, filas or [])))

    ayuda = libro.create_sheet("Instrucciones")
    for fila in instrucciones:
        ayuda.append(list(fila))
    for celda in ayuda[1] if instrucciones else []:
        celda.font = Font(bold=True)
    for indice in range(1, max((len(f) for f in instrucciones), default=0) + 1):
        ayuda.column_dimensions[get_column_letter(indice)].width = 40
    _textos_como_texto(ayuda)

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def generar_reporte(titulo: str, filtros: list[tuple[str, str]], columnas: list[tuple[str, str | None]],
                    filas: list[list], totales: list) -> bytes:
    """Libro de un reporte (módulo 09): título, filtros aplicados, una fila en
    blanco, encabezados en negrita, filas, fila de totales en negrita, formato
    numérico por columna (`columnas` = [(título, formato de Excel o None)]),
    anchos ajustados y panel congelado bajo el encabezado. Protegido contra
    fórmulas con `_textos_como_texto`."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    hoja.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=14)
    fila = 2
    for etiqueta, valor in filtros:
        hoja.cell(row=fila, column=1, value=f"{etiqueta}: {valor}")
        fila += 1
    encabezado = fila + 1   # una fila en blanco entre los filtros y la tabla
    for indice, (nombre, _) in enumerate(columnas, start=1):
        hoja.cell(row=encabezado, column=indice, value=nombre).font = Font(bold=True)
    for desplazamiento, valores in enumerate([*filas, totales], start=1):
        for indice, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=encabezado + desplazamiento, column=indice, value=valor)
            formato = columnas[indice - 1][1]
            if formato:
                celda.number_format = formato
    for celda in hoja[encabezado + len(filas) + 1]:
        celda.font = Font(bold=True)
    for indice, (nombre, _) in enumerate(columnas, start=1):
        valores = [nombre, *(f[indice - 1] for f in filas), totales[indice - 1]]
        largo = max(len(str(v)) for v in valores if v is not None)
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(12, largo + 4), 60)
    hoja.freeze_panes = f"A{encabezado + 1}"
    _textos_como_texto(hoja)

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def celda_decimal(valor) -> Decimal | None:
    """Número de una celda. Acepta números y texto con coma o punto decimal
    ("1.234,56", "1234,5", "1234.5"). None si está vacía; ValueError si no es
    un número finito."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        raise ValueError("no es un número")
    if isinstance(valor, (int, Decimal)):
        numero = Decimal(valor)
    elif isinstance(valor, float):
        numero = Decimal(str(valor))
    else:
        texto = str(valor).strip().replace(" ", "")
        if not texto:
            return None
        if "," in texto and "." in texto:
            texto = texto.replace(".", "").replace(",", ".")
        elif "," in texto:
            texto = texto.replace(",", ".")
        try:
            numero = Decimal(texto)
        except InvalidOperation as exc:
            raise ValueError("no es un número") from exc
    if not numero.is_finite():
        raise ValueError("no es un número")
    return numero


def celda_entero(valor) -> int | None:
    numero = celda_decimal(valor)
    if numero is None:
        return None
    if numero != numero.to_integral_value():
        raise ValueError("no es un número entero")
    return int(numero)
