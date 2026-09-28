from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.errores import DatoInvalido, NoEncontrado, ReglaDeNegocio
from app.models import (
    Auditoria, EstadoAuditoriaInventario, MotivoBaja, MovimientoStock, TipoMovimientoStock,
)
from app.reports import auditoria_pdf
from app.schemas.categorias import CategoriaCrear
from app.schemas.productos import ProductoCrear
from app.schemas.proveedores import ProveedorCrear
from app.schemas.stock import AuditoriaIniciar, ItemContado
from app.services import auditoria_inventario as servicio
from app.services import categorias as servicio_categorias
from app.services import productos as servicio_productos
from app.services import proveedores as servicio_proveedores
from app.services import stock as servicio_stock


@pytest.fixture
def base(db, crear_usuario):
    admin = crear_usuario("admin")
    ropa = servicio_categorias.crear_categoria(db, admin, CategoriaCrear(nombre="Ropa auditoria"), None)
    hoja = servicio_categorias.crear_categoria(
        db, admin, CategoriaCrear(nombre="Remeras auditoria", parent_id=ropa.id), None
    )
    prov = servicio_proveedores.crear_proveedor(db, admin, ProveedorCrear(nombre="Prov auditoria", moneda="ars"), None)

    def alta(descripcion, **extra):
        datos = {"categoria_id": hoja.id, "proveedor_id": prov.id, "costo": "1000", "margen": "0",
                 "descripcion": descripcion}
        datos.update(extra)
        return servicio_productos.crear_producto(db, admin, ProductoCrear(**datos), None)

    return {"admin": admin, "ropa": ropa, "hoja": hoja, "prov": prov, "alta": alta,
            "motivo": db.scalar(select(MotivoBaja).limit(1))}


# --- 1. iniciar --------------------------------------------------------

def test_iniciar_crea_en_curso_y_audita(db, base):
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas="Conteo anual"), None)
    assert auditoria.estado == EstadoAuditoriaInventario.EN_CURSO
    assert auditoria.usuario_id == base["admin"].id
    assert auditoria.notas == "Conteo anual"
    registro = db.scalar(
        select(Auditoria).where(Auditoria.accion == "auditoria_inventario.iniciar", Auditoria.entidad_id == auditoria.id)
    )
    assert registro is not None


def test_iniciar_dos_veces_da_409(db, base):
    servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    with pytest.raises(ReglaDeNegocio) as error:
        servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    assert error.value.mensaje == "Ya hay un conteo en curso."


def test_iniciar_categoria_inexistente(db, base):
    with pytest.raises(NoEncontrado):
        servicio.iniciar(db, base["admin"], AuditoriaIniciar(filtro_categoria_id=999999), None)


# --- 2. registrar_items -------------------------------------------------

def test_registrar_items_captura_stock_al_registrar(db, base):
    p = base["alta"]("Remera conteo A")
    servicio_stock.registrar_ingreso(db, base["admin"], p.id, 10, None, None)
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    items = servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=8)], None)
    assert len(items) == 1
    item = items[0]
    assert item.cantidad_sistema == 10 and item.cantidad_contada == 8 and item.diferencia == -2


def test_registrar_items_de_nuevo_sobreescribe_y_recaptura_sistema(db, base):
    p = base["alta"]("Remera conteo B")
    servicio_stock.registrar_ingreso(db, base["admin"], p.id, 10, None, None)
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=8)], None)
    # El stock cambia entre un conteo y otro (ingreso adicional).
    servicio_stock.registrar_ingreso(db, base["admin"], p.id, 5, None, None)
    items = servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=12)], None)
    assert len(items) == 1
    item = items[0]
    assert item.cantidad_sistema == 15 and item.cantidad_contada == 12 and item.diferencia == -3
    # Sigue habiendo un solo ítem para ese producto (upsert, no duplicado).
    total_items = db.scalars(select(servicio.AuditoriaItem).where(servicio.AuditoriaItem.auditoria_id == auditoria.id)).all()
    assert len(total_items) == 1


def test_registrar_items_producto_infinito(db, base):
    infinito = base["alta"]("Remera conteo infinita", stock_infinito=True)
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    with pytest.raises(DatoInvalido):
        servicio.registrar_items(
            db, base["admin"], auditoria, [ItemContado(producto_id=infinito.id, cantidad_contada=1)], None
        )


def test_registrar_items_producto_inexistente(db, base):
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    with pytest.raises(NoEncontrado):
        servicio.registrar_items(
            db, base["admin"], auditoria, [ItemContado(producto_id=999999, cantidad_contada=1)], None
        )


def test_registrar_items_auditoria_cerrada(db, base):
    p = base["alta"]("Remera conteo cerrada")
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=0)], None)
    servicio.finalizar(db, base["admin"], auditoria, None)
    with pytest.raises(ReglaDeNegocio):
        servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=1)], None)


# --- 3. editar_item / eliminar_item -------------------------------------

def test_editar_item_solo_en_curso(db, base):
    p = base["alta"]("Remera editar item")
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    items = servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=3)], None)
    item = servicio.editar_item(db, base["admin"], auditoria, items[0].id, 5, None)
    assert item.cantidad_contada == 5
    servicio.finalizar(db, base["admin"], auditoria, None)
    with pytest.raises(ReglaDeNegocio):
        servicio.editar_item(db, base["admin"], auditoria, items[0].id, 6, None)


def test_editar_item_de_otra_auditoria_da_404(db, base):
    p = base["alta"]("Remera item ajeno")
    auditoria1 = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    items = servicio.registrar_items(db, base["admin"], auditoria1, [ItemContado(producto_id=p.id, cantidad_contada=1)], None)
    servicio.finalizar(db, base["admin"], auditoria1, None)
    auditoria2 = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    with pytest.raises(NoEncontrado):
        servicio.editar_item(db, base["admin"], auditoria2, items[0].id, 2, None)


def test_eliminar_item_solo_en_curso(db, base):
    p = base["alta"]("Remera eliminar item")
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    items = servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=3)], None)
    item_id = items[0].id
    assert [i.id for i in auditoria.items] == [item_id]          # colección ya cargada
    servicio.eliminar_item(db, base["admin"], auditoria, item_id, None)
    assert db.get(servicio.AuditoriaItem, item_id) is None
    assert auditoria.items == []                                 # la misma instancia ya no lo muestra


def test_eliminar_item_de_otra_auditoria_da_404(db, base):
    p = base["alta"]("Remera item ajeno 2")
    auditoria1 = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    items = servicio.registrar_items(db, base["admin"], auditoria1, [ItemContado(producto_id=p.id, cantidad_contada=1)], None)
    servicio.finalizar(db, base["admin"], auditoria1, None)
    auditoria2 = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    with pytest.raises(NoEncontrado):
        servicio.eliminar_item(db, base["admin"], auditoria2, items[0].id, None)


# --- 4. finalizar --------------------------------------------------------

def test_finalizar_sin_items_da_409(db, base):
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    with pytest.raises(ReglaDeNegocio) as error:
        servicio.finalizar(db, base["admin"], auditoria, None)
    assert error.value.mensaje == "Registrá al menos un producto antes de cerrar el conteo."


def test_finalizar_genera_ajustes_por_diferencia(db, base):
    a = base["alta"]("Remera cierre A")
    b = base["alta"]("Remera cierre B (sin diferencia)")
    servicio_stock.registrar_ingreso(db, base["admin"], a.id, 10, None, None)
    servicio_stock.registrar_ingreso(db, base["admin"], b.id, 5, None, None)
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    servicio.registrar_items(
        db, base["admin"], auditoria,
        [ItemContado(producto_id=a.id, cantidad_contada=7), ItemContado(producto_id=b.id, cantidad_contada=5)],
        None,
    )
    cerrada = servicio.finalizar(db, base["admin"], auditoria, None)
    assert cerrada.estado == EstadoAuditoriaInventario.CERRADA
    assert cerrada.fecha_fin is not None
    assert a.stock == 7          # el ajuste dejó el stock en lo contado
    assert b.stock == 5          # sin diferencia: no hubo ajuste

    ajustes = db.scalars(
        select(MovimientoStock).where(
            MovimientoStock.tipo == TipoMovimientoStock.AJUSTE_AUDITORIA, MovimientoStock.auditoria_id == auditoria.id
        )
    ).all()
    assert len(ajustes) == 1 and ajustes[0].producto_id == a.id and ajustes[0].cantidad == -3

    registro = db.scalar(
        select(Auditoria).where(Auditoria.accion == "auditoria_inventario.cerrar", Auditoria.entidad_id == auditoria.id)
    )
    assert registro is not None

    # Después de cerrar se puede iniciar otra.
    otra = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    assert otra.estado == EstadoAuditoriaInventario.EN_CURSO


def test_finalizar_con_faltante_no_aplica_nada_y_sigue_en_curso(db, base):
    """Se vendió entre el conteo y el cierre: al cerrar, el ajuste dejaría el
    stock negativo (aplicar_movimiento no permite faltante en un ajuste de
    auditoría) -> 409, la auditoría sigue en curso y no queda ningún ajuste."""
    p = base["alta"]("Remera faltante cierre")
    servicio_stock.registrar_ingreso(db, base["admin"], p.id, 10, None, None)
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=8)], None)
    # Venta con permitir_faltante después de registrar el conteo: baja mucho más el stock real.
    servicio_stock.aplicar_movimiento(
        db, base["admin"], tipo=TipoMovimientoStock.VENTA, producto_id=p.id, cantidad=-9,
        referencia_venta_id=1, permitir_faltante=True,
    )
    assert p.stock == 1
    with pytest.raises(ReglaDeNegocio):
        servicio.finalizar(db, base["admin"], auditoria, None)
    db.refresh(auditoria)
    assert auditoria.estado == EstadoAuditoriaInventario.EN_CURSO
    assert p.stock == 1          # no se aplicó ningún ajuste
    assert db.scalar(
        select(MovimientoStock).where(
            MovimientoStock.tipo == TipoMovimientoStock.AJUSTE_AUDITORIA, MovimientoStock.auditoria_id == auditoria.id
        )
    ) is None


# --- 6. listar ------------------------------------------------------------

def test_listar_por_estado_y_fechas(db, base):
    p = base["alta"]("Remera listar")
    en_curso = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    servicio.registrar_items(db, base["admin"], en_curso, [ItemContado(producto_id=p.id, cantidad_contada=0)], None)
    servicio.finalizar(db, base["admin"], en_curso, None)
    otra_en_curso = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)

    cerradas, total_cerradas = servicio.listar(db, estado=EstadoAuditoriaInventario.CERRADA)
    assert total_cerradas >= 1 and all(a.estado == EstadoAuditoriaInventario.CERRADA for a in cerradas)

    en_curso_lista, total_en_curso = servicio.listar(db, estado=EstadoAuditoriaInventario.EN_CURSO)
    assert {a.id for a in en_curso_lista} == {otra_en_curso.id}
    assert total_en_curso == 1

    todas, total_todas = servicio.listar(db, desde=date.today(), hasta=date.today())
    assert total_todas >= 2


# --- 7. PDF y Excel --------------------------------------------------------

def test_generar_pdf_auditoria(db, base):
    p = base["alta"]("Remera PDF")
    servicio_stock.registrar_ingreso(db, base["admin"], p.id, 10, None, None)
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas="Conteo con PDF"), None)
    servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=7)], None)

    contenido_html = auditoria_pdf.pdf.html(auditoria_pdf.PLANTILLA, {"auditoria": auditoria})
    assert p.codigo_con_verificador in contenido_html
    assert "-3" in contenido_html

    pdf_bytes = auditoria_pdf.generar_pdf_auditoria(db, auditoria)
    assert pdf_bytes.startswith(b"%PDF")


def test_generar_xls_auditoria(db, base):
    a = base["alta"]("Remera excel")
    servicio_stock.registrar_ingreso(db, base["admin"], a.id, 10, None, None)
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=a.id, cantidad_contada=6)], None)

    contenido = auditoria_pdf.generar_xls_auditoria(db, auditoria)
    libro = load_workbook(BytesIO(contenido))
    hoja = libro.active
    encabezados = [c.value for c in hoja[1]]
    assert encabezados == ["Código", "Descripción", "Sistema", "Contado", "Diferencia"]
    filas = list(hoja.iter_rows(min_row=2, values_only=True))
    assert len(filas) == 1
    assert filas[0] == (a.codigo_con_verificador, a.descripcion, 10, 6, -4)


def test_generar_xls_auditoria_descripcion_con_formula_queda_como_texto(db, base):
    p = base["alta"]("=HYPERLINK(\"http://x\",\"x\")")
    auditoria = servicio.iniciar(db, base["admin"], AuditoriaIniciar(notas=None), None)
    servicio.registrar_items(db, base["admin"], auditoria, [ItemContado(producto_id=p.id, cantidad_contada=0)], None)

    contenido = auditoria_pdf.generar_xls_auditoria(db, auditoria)
    libro = load_workbook(BytesIO(contenido))
    hoja = libro.active
    celda_descripcion = hoja.cell(row=2, column=2)
    assert celda_descripcion.data_type == "s"
    assert celda_descripcion.value == p.descripcion
