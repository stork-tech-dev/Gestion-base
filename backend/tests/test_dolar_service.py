import sys
import types
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.errores import DatoInvalido, ReglaDeNegocio
from app.models import Auditoria, EstadoProveedor, OrigenCambioDolar, ProveedorDolarHistorial
from app.schemas.proveedores import ModalidadMasiva, ProveedorCrear
from app.services import proveedores as servicio


@pytest.fixture
def admin(crear_usuario):
    return crear_usuario("admin")


def _prov(db, admin, nombre, moneda="usd", dolar="1000"):
    datos = {"nombre": nombre, "moneda": moneda}
    if moneda == "usd":
        datos["dolar_actual"] = dolar
    return servicio.crear_proveedor(db, admin, ProveedorCrear(**datos), None)


def _historial(db, proveedor):
    return db.scalars(select(ProveedorDolarHistorial).where(ProveedorDolarHistorial.proveedor_id == proveedor.id)).all()


def test_cambio_individual_registra_historial_y_auditoria(db, admin, monkeypatch):
    recalculados = []
    monkeypatch.setattr(servicio, "_recalcular_precios", lambda db, pid: recalculados.append(pid))
    p = _prov(db, admin, "Importadora")
    servicio.cambiar_dolar(db, admin, p, Decimal("1100.005"), None)
    assert p.dolar_actual == Decimal("1100.01")
    (fila,) = _historial(db, p)
    assert (fila.valor_anterior, fila.valor_nuevo, fila.origen) == (Decimal("1000"), Decimal("1100.01"), OrigenCambioDolar.MANUAL)
    registro = db.scalar(select(Auditoria).where(Auditoria.accion == "dolar.cambio", Auditoria.entidad_id == p.id))
    assert registro.estado_anterior == {"dolar_actual": "1000.00"} and registro.estado_nuevo["origen"] == "manual"
    assert recalculados == [p.id]


def test_mismo_valor_no_genera_nada(db, admin, monkeypatch):
    recalculados = []
    monkeypatch.setattr(servicio, "_recalcular_precios", lambda db, pid: recalculados.append(pid))
    p = _prov(db, admin, "Sin cambios")
    servicio.cambiar_dolar(db, admin, p, Decimal("1000.00"), None)
    assert _historial(db, p) == [] and recalculados == []
    assert db.scalar(select(func.count()).select_from(Auditoria).where(Auditoria.accion == "dolar.cambio", Auditoria.entidad_id == p.id)) == 0


def test_dolar_de_proveedor_en_pesos_da_409(db, admin):
    p = _prov(db, admin, "En pesos", moneda="ars")
    with pytest.raises(ReglaDeNegocio) as error:
        servicio.cambiar_dolar(db, admin, p, Decimal("100"), None)
    assert error.value.mensaje == "El proveedor trabaja en pesos: no tiene valor de dólar."


def test_dolar_no_positivo(db, admin):
    p = _prov(db, admin, "Cero")
    with pytest.raises(DatoInvalido):
        servicio.cambiar_dolar(db, admin, p, Decimal("0"), None)


def test_historial_mas_reciente_primero(db, admin):
    p = _prov(db, admin, "Con historial")
    servicio.cambiar_dolar(db, admin, p, Decimal("1100"), None)
    servicio.cambiar_dolar(db, admin, p, Decimal("1200"), None)
    historial = servicio.historial_dolar(db, p.id)
    assert [h["valor_nuevo"] for h in historial] == [Decimal("1200.00"), Decimal("1100.00")]
    assert historial[0]["usuario_nombre"] == admin.nombre


def test_masivo_por_valor_y_porcentaje(db, admin):
    a = _prov(db, admin, "A dólares", dolar="1000")
    b = _prov(db, admin, "B dólares", dolar="1500")
    pesos = _prov(db, admin, "C pesos", moneda="ars")
    inactivo = _prov(db, admin, "D inactivo")
    servicio.cambiar_estado(db, admin, inactivo, EstadoProveedor.DESACTIVADO, False, None)

    vista = servicio.preview_masivo(db, None, ModalidadMasiva.PORCENTAJE, Decimal("10"))
    assert [(v["nombre"], v["valor_nuevo"]) for v in vista] == [("A dólares", Decimal("1100.00")), ("B dólares", Decimal("1650.00"))]
    assert a.dolar_actual == Decimal("1000")  # la vista previa no aplica nada

    servicio.aplicar_masivo(db, admin, [a.id, pesos.id, inactivo.id], ModalidadMasiva.VALOR, Decimal("1300"), None)
    assert a.dolar_actual == Decimal("1300.00") and b.dolar_actual == Decimal("1500")
    assert _historial(db, a)[0].origen is OrigenCambioDolar.MASIVO_VALOR

    with pytest.raises(ReglaDeNegocio):
        servicio.aplicar_masivo(db, admin, [pesos.id, inactivo.id], ModalidadMasiva.VALOR, Decimal("1300"), None)
    with pytest.raises(DatoInvalido):
        servicio.preview_masivo(db, None, ModalidadMasiva.PORCENTAJE, Decimal("-100"))
    with pytest.raises(DatoInvalido):
        servicio.preview_masivo(db, None, ModalidadMasiva.VALOR, Decimal("0"))


def test_plantilla_solo_activos_en_dolares(db, admin):
    a = _prov(db, admin, "Plantilla A", dolar="1000")
    _prov(db, admin, "Plantilla pesos", moneda="ars")
    libro = load_workbook(BytesIO(servicio.generar_plantilla_dolar(db)))
    hoja = libro["Dolar"]
    assert [c.value for c in hoja[1]] == ["proveedor_id", "nombre", "dolar_nuevo"]
    filas = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=2)]
    assert [a.id, "Plantilla A", 1000] in [[f[0], f[1], int(f[2])] for f in filas]
    assert all(f[1] != "Plantilla pesos" for f in filas)
    assert "Instrucciones" in libro.sheetnames


def _xlsx(filas):
    libro = Workbook()
    for fila in filas:
        libro.active.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def test_importar_todo_o_nada(db, admin):
    a = _prov(db, admin, "Excel A", dolar="1000")
    b = _prov(db, admin, "Excel B", dolar="1000")
    pesos = _prov(db, admin, "Excel pesos", moneda="ars")
    contenido = _xlsx([
        ["proveedor_id", "dolar_nuevo"],
        [a.id, 1100],
        [b.id, "abc"],
        [a.id, 1200],
        [pesos.id, 900],
        [999999, 900],
        ["x", 900],
    ])
    resultado = servicio.importar_dolar(db, admin, contenido, None)
    assert resultado["aplicados"] == 0 and resultado["sin_cambios"] == 0
    assert [e["fila"] for e in resultado["errores"]] == [3, 4, 5, 6, 7]
    assert a.dolar_actual == Decimal("1000")  # no se aplicó nada


def test_importar_aplica_y_cuenta_sin_cambios(db, admin):
    a = _prov(db, admin, "Excel OK A", dolar="1000")
    b = _prov(db, admin, "Excel OK B", dolar="1000")
    resultado = servicio.importar_dolar(db, admin, _xlsx([["dolar_nuevo", "proveedor_id"], ["1.150,50", a.id], [1000, b.id]]), None)
    assert resultado == {"aplicados": 1, "sin_cambios": 1, "errores": []}
    assert a.dolar_actual == Decimal("1150.50")
    assert _historial(db, a)[0].origen is OrigenCambioDolar.IMPORTACION_EXCEL


def test_la_plantilla_se_puede_volver_a_subir_tal_cual(db, admin):
    _prov(db, admin, "Ida y vuelta", dolar="1000")
    resultado = servicio.importar_dolar(db, admin, servicio.generar_plantilla_dolar(db), None)
    assert resultado["errores"] == [] and resultado["aplicados"] == 0


def test_recalcular_precios_sin_modulo_productos_no_hace_nada(db, monkeypatch):
    # El módulo app.services.productos todavía no existe (llega en el módulo 04):
    # nos asegura que no está en sys.modules para que el import realmente falle.
    monkeypatch.delitem(sys.modules, "app.services.productos", raising=False)
    assert servicio._recalcular_precios(db, 1) is None


def test_recalcular_precios_con_modulo_incompleto_no_esconde_el_error(db, monkeypatch):
    # Un módulo app.services.productos que EXISTE pero no tiene la función
    # esperada (ej. le cambiaron el nombre) tiene que romper, no tragarse el error.
    falso = types.ModuleType("app.services.productos")
    monkeypatch.setitem(sys.modules, "app.services.productos", falso)
    with pytest.raises(ImportError):
        servicio._recalcular_precios(db, 1)


def _auditorias_dolar(db, proveedor):
    return db.scalar(select(func.count()).select_from(Auditoria).where(
        Auditoria.accion == "dolar.cambio", Auditoria.entidad_id == proveedor.id))


def test_masivo_porcentaje_enorme_da_dato_invalido(db, admin):
    _prov(db, admin, "Porcentaje enorme", dolar="1000")
    with pytest.raises(DatoInvalido) as error:
        servicio.aplicar_masivo(db, admin, None, ModalidadMasiva.PORCENTAJE, Decimal("9999999"), None)
    assert error.value.mensaje == "El valor del dólar es demasiado grande."


def test_masivo_marca_aplicado_y_registra_solo_los_que_cambian(db, admin, monkeypatch):
    recalculados = []
    monkeypatch.setattr(servicio, "_recalcular_precios", lambda db, pid: recalculados.append(pid))
    cambia = _prov(db, admin, "Masivo cambia", dolar="1000")
    igual = _prov(db, admin, "Masivo igual", dolar="1300")
    vista = servicio.preview_masivo(db, [cambia.id, igual.id], ModalidadMasiva.VALOR, Decimal("1300"))
    assert all(f["aplicado"] is False for f in vista)

    vista = servicio.aplicar_masivo(db, admin, [cambia.id, igual.id], ModalidadMasiva.VALOR, Decimal("1300"), None)
    assert {f["proveedor_id"]: f["aplicado"] for f in vista} == {cambia.id: True, igual.id: False}
    assert len(_historial(db, cambia)) == 1 and _historial(db, igual) == []
    assert _auditorias_dolar(db, cambia) == 1 and _auditorias_dolar(db, igual) == 0
    assert recalculados == [cambia.id]


def test_importar_registra_solo_los_que_cambian_y_audita_la_importacion(db, admin, monkeypatch):
    recalculados = []
    monkeypatch.setattr(servicio, "_recalcular_precios", lambda db, pid: recalculados.append(pid))
    cambia = _prov(db, admin, "Excel cambia", dolar="1000")
    igual = _prov(db, admin, "Excel igual", dolar="1000")
    resultado = servicio.importar_dolar(db, admin, _xlsx([["proveedor_id", "dolar_nuevo"], [cambia.id, 1200], [igual.id, 1000]]), None)
    assert resultado == {"aplicados": 1, "sin_cambios": 1, "errores": []}
    assert len(_historial(db, cambia)) == 1 and _historial(db, igual) == []
    assert _auditorias_dolar(db, cambia) == 1 and _auditorias_dolar(db, igual) == 0
    assert recalculados == [cambia.id]
    (registro,) = db.scalars(select(Auditoria).where(Auditoria.accion == "dolar.importar")).all()
    assert registro.entidad == "proveedor" and registro.entidad_id is None and registro.usuario_id == admin.id
    assert registro.estado_nuevo == {"aplicados": 1, "sin_cambios": 1}


def test_importar_con_errores_no_audita_la_importacion(db, admin):
    _prov(db, admin, "Excel con error", dolar="1000")
    resultado = servicio.importar_dolar(db, admin, _xlsx([["proveedor_id", "dolar_nuevo"], [999999, 10]]), None)
    assert resultado["errores"]
    assert db.scalar(select(func.count()).select_from(Auditoria).where(Auditoria.accion == "dolar.importar")) == 0


@pytest.mark.parametrize("valor, mensaje", [
    (0.001, "El valor del dólar debe ser mayor a cero."),
    (1e12, "El valor del dólar es demasiado grande."),
])
def test_importar_valores_fuera_de_rango_son_error_de_fila(db, admin, valor, mensaje):
    a = _prov(db, admin, "Excel rango A", dolar="1000")
    b = _prov(db, admin, "Excel rango B", dolar="1000")
    resultado = servicio.importar_dolar(db, admin, _xlsx([["proveedor_id", "dolar_nuevo"], [a.id, 1100], [b.id, valor]]), None)
    assert resultado["aplicados"] == 0 and resultado["errores"] == [{"fila": 3, "error": mensaje}]
    assert a.dolar_actual == Decimal("1000") and _historial(db, a) == []


def test_importar_fila_sin_modificar_de_proveedor_desactivado_es_sin_cambios(db, admin):
    """Se bajó el Excel, después se desactivó un proveedor y se sube el Excel
    sin tocar su fila: no es un error, cuenta como sin cambios."""
    activo = _prov(db, admin, "Excel sigue activo", dolar="1000")
    desactivado = _prov(db, admin, "Excel desactivado", dolar="1000")
    contenido = servicio.generar_plantilla_dolar(db)
    servicio.cambiar_estado(db, admin, desactivado, EstadoProveedor.DESACTIVADO, False, None)
    resultado = servicio.importar_dolar(db, admin, contenido, None)
    assert resultado["errores"] == [] and resultado["aplicados"] == 0 and resultado["sin_cambios"] >= 2

    # Si la fila del desactivado SÍ cambia el dólar, sigue siendo un error.
    resultado = servicio.importar_dolar(db, admin, _xlsx([["proveedor_id", "dolar_nuevo"], [activo.id, 1000], [desactivado.id, 1100]]), None)
    assert resultado["errores"] == [{"fila": 3, "error": "El proveedor no está activo"}]
