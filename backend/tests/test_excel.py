from decimal import Decimal
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errores import ReglaDeNegocio
from app.core.excel import (
    FILAS_MAXIMAS,
    celda_decimal,
    celda_entero,
    clave_columna,
    generar_plantilla,
    leer_tabla,
)


def _xlsx(filas) -> bytes:
    libro = Workbook()
    hoja = libro.active
    for fila in filas:
        hoja.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def test_clave_columna():
    assert clave_columna("  Dólar   Nuevo ") == "dolar nuevo"
    assert clave_columna("Proveedor_ID") == "proveedor_id"
    assert clave_columna(None) == ""


def test_leer_tabla_ubica_columnas_por_nombre_y_saltea_vacias():
    contenido = _xlsx([
        ["Extra", "DOLAR_NUEVO", "Proveedor_ID"],
        ["x", 1250.5, 3],
        [None, None, None],
        ["", "  ", None],
        ["y", "1.300,25", "4"],
    ])
    filas = leer_tabla(contenido, ["proveedor_id", "dolar_nuevo"], ["nombre"])
    assert filas == [
        (2, {"proveedor_id": 3, "dolar_nuevo": 1250.5, "nombre": None}),
        (5, {"proveedor_id": "4", "dolar_nuevo": "1.300,25", "nombre": None}),
    ]


@pytest.mark.parametrize(
    "contenido, mensaje",
    [
        (b"", "vacío"),
        (b"no soy un excel", "no es un Excel"),
        (None, "Faltan columnas: dolar_nuevo"),
    ],
)
def test_leer_tabla_rechaza_archivos_invalidos(contenido, mensaje):
    if contenido is None:
        contenido = _xlsx([["proveedor_id"], [1]])
    with pytest.raises(ReglaDeNegocio) as error:
        leer_tabla(contenido, ["proveedor_id", "dolar_nuevo"])
    assert mensaje in error.value.mensaje


def test_leer_tabla_limita_filas():
    contenido = _xlsx([["proveedor_id"]] + [[i] for i in range(FILAS_MAXIMAS + 1)])
    with pytest.raises(ReglaDeNegocio) as error:
        leer_tabla(contenido, ["proveedor_id"])
    assert str(FILAS_MAXIMAS) in error.value.mensaje


def test_plantilla_ida_y_vuelta():
    contenido = generar_plantilla(
        "Dolar", ["proveedor_id", "nombre", "dolar_nuevo"],
        [["Columna", "Qué poner"], ["dolar_nuevo", "El valor nuevo"]],
        filas=[[1, "Importadora Sur", Decimal("1250.50")]],
    )
    libro = load_workbook(BytesIO(contenido))
    assert libro.sheetnames == ["Dolar", "Instrucciones"]
    assert libro["Dolar"]["A1"].font.bold and libro["Dolar"].freeze_panes == "A2"
    assert libro["Instrucciones"]["A1"].value == "Columna"
    filas = leer_tabla(contenido, ["proveedor_id", "dolar_nuevo"], ["nombre"])
    assert filas == [(2, {"proveedor_id": 1, "dolar_nuevo": 1250.5, "nombre": "Importadora Sur"})]


@pytest.mark.parametrize(
    "valor, esperado",
    [
        (None, None), ("", None), ("  ", None), (3, Decimal("3")), (1250.5, Decimal("1250.5")),
        (Decimal("7.25"), Decimal("7.25")), ("1.234,56", Decimal("1234.56")), ("1234,5", Decimal("1234.5")),
        ("1234.5", Decimal("1234.5")), (" 1 250 ", Decimal("1250")),
    ],
)
def test_celda_decimal(valor, esperado):
    assert celda_decimal(valor) == esperado


@pytest.mark.parametrize("valor", ["abc", "NaN", "Infinity", True, "1,2,3"])
def test_celda_decimal_invalida(valor):
    with pytest.raises(ValueError):
        celda_decimal(valor)


def test_celda_entero():
    assert celda_entero(12) == 12 and celda_entero(12.0) == 12 and celda_entero("7") == 7
    assert celda_entero(None) is None
    with pytest.raises(ValueError):
        celda_entero(3.5)


def test_leer_tabla_rechaza_xml_truncado():
    """Archivo xlsx válido como zip pero con XML truncado en la hoja."""
    contenido_valido = _xlsx([["proveedor_id", "dolar_nuevo"], [1, 1250.5]])
    buffer = BytesIO()
    with ZipFile(BytesIO(contenido_valido), "r") as zip_entrada:
        with ZipFile(buffer, "w") as zip_salida:
            for item in zip_entrada.infolist():
                if item.filename == "xl/worksheets/sheet1.xml":
                    # Truncar el contenido del XML
                    contenido = zip_entrada.read(item.filename)
                    zip_salida.writestr(item, contenido[:len(contenido) // 2])
                else:
                    zip_salida.writestr(item, zip_entrada.read(item.filename))
    corrupto = buffer.getvalue()
    with pytest.raises(ReglaDeNegocio) as error:
        leer_tabla(corrupto, ["proveedor_id", "dolar_nuevo"])
    assert "no es un Excel" in error.value.mensaje


def test_leer_tabla_rechaza_xml_corrupto_despues_encabezado():
    """Archivo xlsx válido hasta el encabezado pero con XML corrupto después."""
    contenido_valido = _xlsx([["proveedor_id", "dolar_nuevo"], [1, 1250.5], [2, 1300.0]])
    buffer = BytesIO()
    with ZipFile(BytesIO(contenido_valido), "r") as zip_entrada:
        with ZipFile(buffer, "w") as zip_salida:
            for item in zip_entrada.infolist():
                if item.filename == "xl/worksheets/sheet1.xml":
                    # Reemplazar contenido con XML que tiene encabezado pero fallo después
                    contenido_corrupto = b'<?xml version="1.0"?><worksheet><sheetData><row r="1"><c r="A1"><v>proveedor_id</v></c></row><row r="2" corrupto'
                    zip_salida.writestr(item, contenido_corrupto)
                else:
                    zip_salida.writestr(item, zip_entrada.read(item.filename))
    corrupto = buffer.getvalue()
    with pytest.raises(ReglaDeNegocio) as error:
        leer_tabla(corrupto, ["proveedor_id", "dolar_nuevo"])
    assert "no es un Excel" in error.value.mensaje


def test_plantilla_no_deja_inyectar_formulas():
    """Un texto que empieza con = + - @ tab o CR se guarda como texto, no
    como fórmula, y leer_tabla lo devuelve igual."""
    nombres = ['=HYPERLINK("http://x","x")', "+1+1", "-2+3", "@SUMA(A1)", "\tTab", "\rRetorno", "Normal"]
    contenido = generar_plantilla("Datos", ["id", "nombre"], [["Ayuda"], ["=1+1"]], [[i, n] for i, n in enumerate(nombres, 1)])
    libro = load_workbook(BytesIO(contenido))
    celdas = [fila[1] for fila in libro["Datos"].iter_rows(min_row=2)]
    assert all(c.data_type == "s" for c in celdas)
    assert celdas[0].value == '=HYPERLINK("http://x","x")'
    assert libro["Instrucciones"]["A2"].data_type == "s"
    assert [f["nombre"] for _, f in leer_tabla(contenido, ["id", "nombre"])][0] == '=HYPERLINK("http://x","x")'


def test_leer_tabla_rechaza_archivo_que_descomprimido_es_enorme():
    # 60 MB de ceros comprimen a unos pocos KB: pasa el límite de 5 MB del
    # archivo, pero no el de lo descomprimido (zip bomb).
    salida = BytesIO()
    with ZipFile(salida, "w", ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (60 * 1024 * 1024))
    with pytest.raises(ReglaDeNegocio, match="descomprimido"):
        leer_tabla(salida.getvalue(), ["nombre"])


def test_leer_tabla_rechaza_zip_con_directorio_central_corrupto():
    # Un .xlsx corrupto puede hacer que ZipFile()/infolist() levante otras
    # excepciones además de BadZipFile (ej. NotImplementedError por una
    # "versión" de zip fuera de rango, en el byte extract_version del
    # directorio central): tienen que dar 409, no un 500.
    contenido = bytearray(_xlsx([["nombre"], ["x"]]))
    firma = b"PK\x01\x02"  # firma de cada registro del directorio central
    posicion = contenido.find(firma)
    assert posicion != -1
    contenido[posicion + 6] = 255  # extract_version: fuera del máximo soportado
    with pytest.raises(ReglaDeNegocio) as error:
        leer_tabla(bytes(contenido), ["nombre"])
    assert "no es un Excel" in error.value.mensaje
