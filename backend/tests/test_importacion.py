"""Importación inicial desde Excel (sesión 08): motor, proveedores y clientes.
Los productos se prueban en la segunda mitad del archivo (Task 2)."""
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.excel import TIPO_XLSX
from app.models import Auditoria, Categoria, Cliente, MovimientoStock, Producto, Proveedor
from app.services.productos import calcular_precio_venta
from tests.fabricas import catalogo

URL = "/api/v1/importacion"
ENCABEZADO_PROVEEDORES = ["nombre", "moneda", "dolar_actual", "pais", "provincia", "telefono", "email", "notas"]
ENCABEZADO_CLIENTES = ["nombre", "dni", "cuit", "condicion_iva", "domicilio", "codigo_postal", "localidad",
                       "telefono", "email"]
CUIT_VALIDO = "20-39059158-7"


def _xlsx(filas) -> bytes:
    libro = Workbook()
    for fila in filas:
        libro.active.append(list(fila))
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _subir(client, h, tipo, accion, filas, nombre="datos.xlsx"):
    return client.post(f"{URL}/{tipo}/{accion}", files={"archivo": (nombre, _xlsx(filas), TIPO_XLSX)}, headers=h)


def _contar(db, modelo):
    return db.scalar(select(func.count()).select_from(modelo))


def _admin(crear_usuario, headers_de):
    return headers_de(crear_usuario("admin"))


def test_plantilla_vacia_da_409_y_con_filas_importa(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    r = client.get(f"{URL}/proveedores/plantilla", headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == TIPO_XLSX
    assert 'filename="importar-proveedores-' in r.headers["content-disposition"]
    libro = load_workbook(BytesIO(r.content))
    assert [c.value for c in libro.worksheets[0][1]] == ENCABEZADO_PROVEEDORES
    assert libro.sheetnames == ["Proveedores", "Instrucciones"]

    r = client.post(f"{URL}/proveedores/importar", files={"archivo": ("p.xlsx", r.content, TIPO_XLSX)}, headers=h)
    assert r.status_code == 409 and r.json()["detail"] == "El archivo no tiene filas para importar."

    libro.worksheets[0].append(["Textil Norte", "Pesos"])
    salida = BytesIO()
    libro.save(salida)
    r = client.post(f"{URL}/proveedores/importar", files={"archivo": ("p.xlsx", salida.getvalue(), TIPO_XLSX)},
                    headers=h)
    assert r.status_code == 200
    assert r.json() == {"tipo": "proveedores", "filas_leidas": 1, "filas_validas": 1, "creados": 1,
                        "categorias_nuevas": [], "errores": [], "aplicado": True}


def test_columnas_en_otro_orden_y_extra(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    filas = [["Columna Extra", "MONEDA", "Nombre", "Dólar Actual"],
             ["x", "Dólares", "Importadora Sur", "1050,5"]]
    r = _subir(client, h, "proveedores", "importar", filas)
    assert r.status_code == 200 and r.json()["creados"] == 1
    proveedor = db.scalar(select(Proveedor).where(Proveedor.nombre == "Importadora Sur"))
    assert proveedor.moneda.value == "usd" and str(proveedor.dolar_actual) == "1050.50"


def test_proveedor_en_dolares_sin_dolar_es_error_de_la_columna(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    r = _subir(client, h, "proveedores", "validar", [["nombre", "moneda"], ["Sin dolar", "usd"]])
    assert r.status_code == 200
    assert r.json()["errores"] == [
        {"fila": 2, "columna": "dolar_actual", "error": "El valor del dólar debe ser mayor a cero."}]


def test_proveedor_en_pesos_ignora_el_dolar(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    r = _subir(client, h, "proveedores", "importar", [["nombre", "moneda", "dolar_actual"], ["En pesos", "ars", "-3"]])
    assert r.json()["aplicado"] is True
    assert db.scalar(select(Proveedor).where(Proveedor.nombre == "En pesos")).dolar_actual is None


def test_proveedor_duplicado_en_base_y_en_archivo(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    assert _subir(client, h, "proveedores", "importar", [["nombre", "moneda"], ["Lanas Río", "Pesos"]]).status_code == 200
    filas = [["nombre", "moneda"], ["lanas rio", "Pesos"], ["Nuevo", "Pesos"], ["NUEVO", "Pesos"]]
    errores = _subir(client, h, "proveedores", "validar", filas).json()["errores"]
    assert errores[0]["fila"] == 2 and errores[0]["columna"] is None and "Ya existe un proveedor" in errores[0]["error"]
    assert errores[1] == {"fila": 4, "columna": "nombre", "error": "Repetido en el archivo (fila 3)."}


def test_moneda_invalida_y_varios_errores_de_una_fila(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    filas = [["nombre", "moneda", "email"], ["x" * 201, "Euros", "no-es-mail"]]
    errores = _subir(client, h, "proveedores", "validar", filas).json()["errores"]
    assert {"fila": 2, "columna": "moneda", "error": "Valor no válido: usá Pesos / Dólares."} in errores
    # moneda falla en el parseo: el alta de esa fila no se intenta, así que no hay error de schema todavía.
    assert len(errores) == 1
    filas = [["nombre", "moneda", "email"], ["x" * 201, "Pesos", "no-es-mail"]]
    errores = _subir(client, h, "proveedores", "validar", filas).json()["errores"]
    assert {"fila": 2, "columna": "nombre", "error": "Admite hasta 200 caracteres."} in errores
    assert {"fila": 2, "columna": "email", "error": "El email no es válido."} in errores


def test_obligatoria_vacia(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    errores = _subir(client, h, "proveedores", "validar", [["nombre", "moneda"], ["", "Pesos"], ["Ok", None]]).json()["errores"]
    assert errores == [{"fila": 2, "columna": "nombre", "error": "Es obligatoria."},
                       {"fila": 3, "columna": "moneda", "error": "Es obligatoria."}]


def test_cliente_cuit_invalido_y_dni_repetido(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    assert _subir(client, h, "clientes", "importar", [["nombre", "dni"], ["Ana", 30123456]]).json()["aplicado"]
    filas = [["nombre", "dni", "cuit", "condicion_iva"],
             ["Beto", "30.123.456", None, None],                 # DNI ya en la base
             ["Caro", None, "20-39059158-2", None],                # dígito verificador mal
             ["Dani", 40111222, None, None],
             ["Eli", "40111222", None, None],                    # DNI repetido en el archivo
             ["Fede", None, None, "Responsable inscripto"]]      # sin CUIT y no es consumidor final
    errores = _subir(client, h, "clientes", "validar", filas).json()["errores"]
    assert [(e["fila"], e["columna"]) for e in errores] == [(2, None), (3, None), (5, "dni"), (6, None)]
    assert "ya está cargado" in errores[0]["error"]
    assert "dígito verificador" in errores[1]["error"]
    assert errores[2]["error"] == "Repetido en el archivo (fila 4)."
    assert "CUIT es obligatorio" in errores[3]["error"]


def test_cliente_condicion_por_etiqueta_valor_o_default(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    filas = [["nombre", "cuit", "condicion_iva"],
             ["Uno SA", CUIT_VALIDO, "responsable_inscripto"],
             ["Dos", None, None],
             ["Tres", "27-28033514-8", "monotributista"]]
    r = _subir(client, h, "clientes", "importar", filas)
    assert r.json()["errores"] == [] and r.json()["creados"] == 3
    condiciones = {c.nombre: c.condicion_iva.value for c in db.scalars(select(Cliente)).all()}
    assert condiciones["Uno SA"] == "responsable_inscripto" and condiciones["Dos"] == "consumidor_final"


def test_todo_o_nada(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    antes = _contar(db, Proveedor)
    filas = [["nombre", "moneda"], ["Bueno 1", "Pesos"], ["Malo", "usd"], ["Bueno 2", "Pesos"]]
    r = _subir(client, h, "proveedores", "importar", filas)
    assert r.status_code == 200
    cuerpo = r.json()
    assert cuerpo["aplicado"] is False and cuerpo["creados"] == 0
    assert cuerpo["filas_leidas"] == 3 and cuerpo["filas_validas"] == 2 and len(cuerpo["errores"]) == 1
    assert _contar(db, Proveedor) == antes


def test_validar_no_deja_rastros(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    proveedores, auditoria = _contar(db, Proveedor), _contar(db, Auditoria)
    r = _subir(client, h, "proveedores", "validar", [["nombre", "moneda"], ["Solo vista previa", "Pesos"]])
    assert r.json() == {"tipo": "proveedores", "filas_leidas": 1, "filas_validas": 1, "creados": 0,
                        "categorias_nuevas": [], "errores": [], "aplicado": False}
    assert _contar(db, Proveedor) == proveedores
    assert _contar(db, Auditoria) == auditoria


def test_importar_audita_la_importacion_y_cada_alta(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    r = _subir(client, h, "clientes", "importar", [["nombre"], ["Uno"], ["Dos"]], nombre="clientes-viejos.xlsx")
    assert r.json()["creados"] == 2
    registro = db.scalar(select(Auditoria).where(Auditoria.accion == "importacion.clientes"))
    assert registro.estado_nuevo == {"archivo": "clientes-viejos.xlsx", "filas": 2, "creados": 2}
    assert registro.entidad == "importacion" and registro.entidad_id is None
    assert db.scalar(select(func.count()).select_from(Auditoria).where(Auditoria.accion == "cliente.crear")) >= 2


def test_sin_permiso_403_y_tipo_invalido_422(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("vendedor"))
    assert client.get(f"{URL}/clientes/plantilla", headers=h).status_code == 403
    assert _subir(client, h, "clientes", "validar", [["nombre"], ["x"]]).status_code == 403
    assert _subir(client, h, "clientes", "importar", [["nombre"], ["x"]]).status_code == 403
    h = _admin(crear_usuario, headers_de)
    assert client.get(f"{URL}/ventas/plantilla", headers=h).status_code == 422


def test_archivo_invalido_409(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    r = client.post(f"{URL}/clientes/validar", files={"archivo": ("x.xlsx", b"basura", TIPO_XLSX)}, headers=h)
    assert r.status_code == 409
    r = _subir(client, h, "clientes", "validar", [["dni"], ["1"]])
    assert r.status_code == 409 and r.json()["detail"] == "Faltan columnas: nombre"


# ---------- Productos ----------

ENCABEZADO_PRODUCTOS = ["descripcion", "categoria", "proveedor", "costo", "margen", "sku_proveedor",
                        "descuento_producto", "peso_gramos", "temporada", "ubicacion_deposito",
                        "stock_inicial", "stock_minimo"]


def _proveedor(client, h, nombre="Prov Import", moneda="Pesos", dolar=None):
    r = _subir(client, h, "proveedores", "importar", [["nombre", "moneda", "dolar_actual"], [nombre, moneda, dolar]])
    assert r.json()["aplicado"], r.json()
    return nombre


def _categoria(db, camino):
    """Categoría final del camino (None si no existe)."""
    parent = None
    for nombre in camino.split(" > "):
        condicion = Categoria.parent_id.is_(None) if parent is None else Categoria.parent_id == parent.id
        parent = db.scalar(select(Categoria).where(Categoria.nombre == nombre, condicion))
        if parent is None:
            return None
    return parent


def test_producto_crea_categorias_del_camino_y_calcula_precio(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    proveedor = _proveedor(client, h, "Dolarizado", "Dólares", "1000")
    filas = [["descripcion", "categoria", "proveedor", "costo", "margen", "temporada", "sku_proveedor"],
             ["Remera lisa", "Ropa Imp > Remeras > Lisas", "dolarizado", "10,5", "50", "Otoño-Invierno", "ABC-1"],
             ["Remera rayada", "ropa imp > REMERAS > Rayadas", proveedor, 12, None, "primavera_verano", None]]
    r = _subir(client, h, "productos", "validar", filas)
    assert r.json()["errores"] == []
    # Los caminos se informan con los nombres guardados, no con el texto de la celda.
    assert r.json()["categorias_nuevas"] == ["Ropa Imp", "Ropa Imp > Remeras", "Ropa Imp > Remeras > Lisas",
                                             "Ropa Imp > Remeras > Rayadas"]
    assert _categoria(db, "Ropa Imp") is None                    # validar no deja categorías

    r = _subir(client, h, "productos", "importar", filas)
    assert r.json()["aplicado"] is True and r.json()["creados"] == 2
    lisa = db.scalar(select(Producto).where(Producto.descripcion == "Remera lisa"))
    assert lisa.categoria_id == _categoria(db, "Ropa Imp > Remeras > Lisas").id
    assert lisa.sku_proveedor == "ABC-1" and lisa.temporada.value == "otoño_invierno"
    # Mismo precio que el alta manual: lo calcula el service de productos.
    assert lisa.precio_venta == calcular_precio_venta(db, Decimal("10.5"), Decimal("50"), lisa.proveedor)
    rayada = db.scalar(select(Producto).where(Producto.descripcion == "Remera rayada"))
    assert rayada.categoria_id == _categoria(db, "Ropa Imp > Remeras > Rayadas").id   # reusó "Ropa Imp > Remeras"


def test_categorias_nuevas_informa_cada_nivel_una_vez(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    proveedor = _proveedor(client, h)
    filas = [["descripcion", "categoria", "proveedor", "costo"],
             ["A", "Hogar Imp > Mantas", proveedor, 100],
             ["B", "Hogar Imp > Mantas", proveedor, 100],
             ["C", "Hogar Imp > Almohadas", proveedor, 100]]
    assert _subir(client, h, "productos", "validar", filas).json()["categorias_nuevas"] == [
        "Hogar Imp", "Hogar Imp > Mantas", "Hogar Imp > Almohadas"]


def test_categoria_con_subcategorias_o_con_productos(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    admin = crear_usuario("admin")
    datos = catalogo(db, admin)             # "Ropa fab N" > "Remeras fab N", con proveedor "Prov fab N"
    datos["alta"]("Existente")               # la subcategoría ya tiene productos
    madre, hija = datos["categoria"].nombre, datos["subcategoria"].nombre
    filas = [["descripcion", "categoria", "proveedor", "costo"],
             ["X", madre, datos["proveedor"].nombre, 100],
             ["Y", f"{madre} > {hija} > Nueva", datos["proveedor"].nombre, 100]]
    errores = _subir(client, h, "productos", "validar", filas).json()["errores"]
    assert errores[0] == {"fila": 2, "columna": "categoria",
                          "error": f"La categoría '{madre}' tiene subcategorías: indicá el camino hasta el último nivel."}
    assert errores[1]["fila"] == 3 and errores[1]["columna"] == "categoria"   # padre con productos (crear_categoria)


def test_proveedor_inexistente_o_inactivo(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    admin = crear_usuario("admin")
    datos = catalogo(db, admin)
    inactivo = datos["proveedor"]
    from app.models import EstadoProveedor
    inactivo.estado = EstadoProveedor.DESACTIVADO
    db.flush()
    camino = f"{datos['categoria'].nombre} > {datos['subcategoria'].nombre}"
    filas = [["descripcion", "categoria", "proveedor", "costo"],
             ["X", camino, "No existe SA", 100],
             ["Y", camino, inactivo.nombre, 100]]
    errores = _subir(client, h, "productos", "validar", filas).json()["errores"]
    assert errores[0] == {"fila": 2, "columna": "proveedor", "error": "No existe el proveedor 'No existe SA'."}
    assert errores[1]["fila"] == 3 and "no está activo" in errores[1]["error"]


def test_stock_inicial_genera_ingreso_y_minimo(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    proveedor = _proveedor(client, h)
    filas = [["descripcion", "categoria", "proveedor", "costo", "stock_inicial", "stock_minimo"],
             ["Con stock", "Stock Imp", proveedor, 100, 12, 3],
             ["Sin stock", "Stock Imp", proveedor, 100, 0, None]]
    assert _subir(client, h, "productos", "importar", filas).json()["aplicado"] is True
    con = db.scalar(select(Producto).where(Producto.descripcion == "Con stock"))
    assert con.stock == 12 and con.stock_minimo == 3
    movimiento = db.scalar(select(MovimientoStock).where(MovimientoStock.producto_id == con.id))
    assert movimiento.tipo.value == "ingreso_proveedor" and movimiento.cantidad == 12
    assert movimiento.notas == "Importación inicial"
    sin = db.scalar(select(Producto).where(Producto.descripcion == "Sin stock"))
    assert sin.stock == 0 and db.scalar(select(MovimientoStock).where(MovimientoStock.producto_id == sin.id)) is None


def test_stock_invalido_es_error_de_columna(client, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    proveedor = _proveedor(client, h)
    filas = [["descripcion", "categoria", "proveedor", "costo", "stock_inicial", "stock_minimo"],
             ["A", "Stock Mal", proveedor, 100, -1, "2,5"]]
    errores = _subir(client, h, "productos", "validar", filas).json()["errores"]
    assert errores == [
        {"fila": 2, "columna": "stock_inicial", "error": "Tiene que ser un número entero entre 0 y 1.000.000."},
        {"fila": 2, "columna": "stock_minimo", "error": "Tiene que ser un número entero."}]


def test_producto_reglas_del_alta_manual(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    proveedor = _proveedor(client, h)
    filas = [["descripcion", "categoria", "proveedor", "costo", "descuento_producto"],
             ["Taza", "Bazar Imp", proveedor, 0, None],             # costo > 0 (schema)
             ["Plato", "Bazar Imp", proveedor, 100, 99],            # descuento > descuento_maximo (service)
             ["Vaso", "Bazar Imp", proveedor, 100, None],
             ["vaso", "bazar imp", proveedor.upper(), 100, None]]   # repetido en el archivo
    errores = _subir(client, h, "productos", "validar", filas).json()["errores"]
    assert errores[0] == {"fila": 2, "columna": "costo", "error": "Tiene que ser mayor a 0."}
    assert errores[1]["fila"] == 3 and "descuento" in errores[1]["error"]
    assert errores[2] == {"fila": 5, "columna": "descripcion", "error": "Repetido en el archivo (fila 4)."}


def test_producto_todo_o_nada_ni_categorias(client, db, crear_usuario, headers_de):
    h = _admin(crear_usuario, headers_de)
    proveedor = _proveedor(client, h)
    productos = _contar(db, Producto)
    filas = [["descripcion", "categoria", "proveedor", "costo", "stock_inicial"],
             ["Bien", "Todo O Nada > Sub", proveedor, 100, 5],
             ["Mal", "Todo O Nada > Sub", "Inexistente", 100, 5]]
    r = _subir(client, h, "productos", "importar", filas)
    assert r.json()["aplicado"] is False and r.json()["creados"] == 0
    assert _contar(db, Producto) == productos
    assert _categoria(db, "Todo O Nada") is None
    assert db.scalar(select(Auditoria).where(Auditoria.accion == "importacion.productos")) is None


def test_importar_productos_sin_permiso_de_stock_ni_productos(client, db, crear_usuario, headers_de):
    """Decisión de producto: alcanza con IMPORTACION crear (rol a medida sin stock ni productos)."""
    from app.models import Rol, UsuarioPermiso
    h_admin = _admin(crear_usuario, headers_de)
    proveedor = _proveedor(client, h_admin)
    rol = Rol(nombre="solo importa", es_sistema=False)
    db.add(rol)
    db.flush()
    usuario = crear_usuario(rol)
    # Mismo patrón que tests/test_paginas_stock.py (permiso puntual por usuario).
    db.add(UsuarioPermiso(usuario_id=usuario.id, modulo="importacion", puede_crear=True))
    db.flush()
    h = headers_de(usuario)
    filas = [["descripcion", "categoria", "proveedor", "costo", "stock_inicial"],
             ["Solo import", "Perm Imp", proveedor, 100, 4]]
    r = _subir(client, h, "productos", "importar", filas)
    assert r.json()["aplicado"] is True, r.json()
    assert db.scalar(select(Producto).where(Producto.descripcion == "Solo import")).stock == 4
