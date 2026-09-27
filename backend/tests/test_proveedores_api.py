from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.core.excel import TIPO_XLSX
from app.models import UsuarioPermiso

URL = "/api/v1/proveedores"


def _alta(client, h, **datos):
    base = {"nombre": "Proveedor API", "moneda": "ars"}
    base.update(datos)
    return client.post(URL, json=base, headers=h)


def test_permisos(client, crear_usuario, headers_de):
    supervisor = headers_de(crear_usuario("supervisor"))
    vendedor = headers_de(crear_usuario("vendedor"))
    assert client.get(URL, headers=supervisor).status_code == 200      # supervisor: solo ver
    assert _alta(client, supervisor).status_code == 403
    assert client.get(URL, headers=vendedor).status_code == 403
    assert client.get(f"{URL}/dolar/plantilla", headers=supervisor).status_code == 403


def test_permiso_puntual_del_cambio_masivo(client, db, crear_usuario, headers_de):
    """El recurso dolar.cambio_masivo alcanza para plantilla/preview, pero no
    da el permiso general de 'editar' (PUT) ni de 'ver' (listado): el override
    queda anclado al recurso puntual y no habilita las filas generales."""
    admin = headers_de(crear_usuario("admin"))
    p = _alta(client, admin, nombre="Recurso puntual", moneda="usd", dolar_actual="1000").json()
    vendedor = crear_usuario("vendedor")
    db.add(UsuarioPermiso(usuario_id=vendedor.id, modulo="proveedores", recurso="dolar.cambio_masivo", puede_editar=True))
    db.flush()
    h = headers_de(vendedor)
    assert client.get(f"{URL}/dolar/plantilla", headers=h).status_code == 200
    cuerpo = {"proveedor_ids": None, "modalidad": "porcentaje", "valor": "10"}
    assert client.post(f"{URL}/dolar/masivo/preview", json=cuerpo, headers=h).status_code == 200
    assert client.put(f"{URL}/{p['id']}", json={"nombre": "Recurso puntual", "moneda": "usd"}, headers=h).status_code == 403
    assert client.get(URL, headers=h).status_code == 403  # el recurso no habilita el "ver" general del módulo


def test_alta_edicion_y_ficha(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("admin"))
    r = _alta(client, h, nombre="Importadora Sur", moneda="usd", dolar_actual="1250.50", pais="China")
    assert r.status_code == 201
    p = r.json()
    assert p["en_dolares"] and p["activo"] and p["dolar_actual"] == "1250.50"
    assert _alta(client, h, nombre="importadora sur").status_code == 409
    assert _alta(client, h, nombre="Sin dólar", moneda="usd").status_code == 400
    r = client.put(f"{URL}/{p['id']}", json={"nombre": "Importadora Sur", "moneda": "usd", "provincia": "Córdoba"}, headers=h)
    assert r.status_code == 200 and r.json()["provincia"] == "Córdoba"
    assert client.get(f"{URL}/{p['id']}", headers=h).json()["nombre"] == "Importadora Sur"
    assert client.get(f"{URL}/999999", headers=h).status_code == 404


def test_listado_con_filtros(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("admin"))
    _alta(client, h, nombre="Córdoba Hilados", provincia="Córdoba")
    _alta(client, h, nombre="Otro", provincia="Salta")
    nombres = [p["nombre"] for p in client.get(URL, params={"provincia": "cordoba"}, headers=h).json()]
    assert nombres == ["Córdoba Hilados"]


def test_estado_y_dolar(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("admin"))
    p = _alta(client, h, nombre="Dólar API", moneda="usd", dolar_actual="1000").json()
    r = client.patch(f"{URL}/{p['id']}/dolar", json={"valor_nuevo": "1100"}, headers=h)
    assert r.status_code == 200 and r.json()["dolar_actual"] == "1100.00"
    historial = client.get(f"{URL}/{p['id']}/dolar/historial", headers=h).json()
    assert historial[0]["valor_nuevo"] == "1100.00" and historial[0]["origen"] == "manual"
    assert client.get(f"{URL}/dolar/en-uso", headers=h).json() == {"hay_proveedores_en_dolares": True}
    r = client.patch(f"{URL}/{p['id']}/estado", json={"estado": "inhabilitado"}, headers=h)
    assert r.status_code == 200 and r.json()["estado"] == "inhabilitado"
    assert client.get(f"{URL}/dolar/en-uso", headers=h).json() == {"hay_proveedores_en_dolares": False}


def test_masivo(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("admin"))
    p = _alta(client, h, nombre="Masivo API", moneda="usd", dolar_actual="1000").json()
    cuerpo = {"proveedor_ids": None, "modalidad": "porcentaje", "valor": "10"}
    vista = client.post(f"{URL}/dolar/masivo/preview", json=cuerpo, headers=h).json()
    assert {"proveedor_id": p["id"], "nombre": "Masivo API", "valor_actual": "1000.00", "valor_nuevo": "1100.00"} in vista
    assert client.get(f"{URL}/{p['id']}", headers=h).json()["dolar_actual"] == "1000.00"
    assert client.post(f"{URL}/dolar/masivo", json=cuerpo, headers=h).status_code == 200
    assert client.get(f"{URL}/{p['id']}", headers=h).json()["dolar_actual"] == "1100.00"


def test_plantilla_e_importacion(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("admin"))
    p = _alta(client, h, nombre="Excel API", moneda="usd", dolar_actual="1000").json()
    r = client.get(f"{URL}/dolar/plantilla", headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == TIPO_XLSX
    assert "attachment" in r.headers["content-disposition"] and "dolar-proveedores-" in r.headers["content-disposition"]
    libro = load_workbook(BytesIO(r.content))
    hoja = libro["Dolar"]
    fila = next(f for f in hoja.iter_rows(min_row=2) if f[0].value == p["id"])
    fila[2].value = 1234.5
    salida = BytesIO()
    libro.save(salida)
    r = client.post(f"{URL}/dolar/importar", files={"archivo": ("dolar.xlsx", salida.getvalue(), TIPO_XLSX)}, headers=h)
    assert r.status_code == 200 and r.json()["aplicados"] == 1 and r.json()["errores"] == []
    assert client.get(f"{URL}/{p['id']}", headers=h).json()["dolar_actual"] == "1234.50"


def test_importacion_con_errores_no_aplica_nada(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("admin"))
    p = _alta(client, h, nombre="Excel errores", moneda="usd", dolar_actual="1000").json()
    libro = Workbook()
    for fila in (["proveedor_id", "dolar_nuevo"], [p["id"], 1500], [999999, 10]):
        libro.active.append(fila)
    salida = BytesIO()
    libro.save(salida)
    r = client.post(f"{URL}/dolar/importar", files={"archivo": ("dolar.xlsx", salida.getvalue(), TIPO_XLSX)}, headers=h)
    assert r.status_code == 200 and r.json()["aplicados"] == 0
    assert r.json()["errores"] == [{"fila": 3, "error": "El proveedor no existe"}]
    assert client.get(f"{URL}/{p['id']}", headers=h).json()["dolar_actual"] == "1000.00"
    r = client.post(f"{URL}/dolar/importar", files={"archivo": ("x.xlsx", b"basura", TIPO_XLSX)}, headers=h)
    assert r.status_code == 409


def test_no_hay_delete(client, crear_usuario, headers_de):
    h = headers_de(crear_usuario("admin"))
    p = _alta(client, h, nombre="Sin delete").json()
    assert client.delete(f"{URL}/{p['id']}", headers=h).status_code == 405
