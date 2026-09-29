"""Reportes (sesión 09), Task 1: motor (filtros, rango, 404/422/403),
exportación a Excel y PDF con auditoría, opciones de filtros y el reporte de
stock valorizado. Los reportes de ventas: test_reportes_ventas.py."""
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.errores import NoProcesable, ReglaDeNegocio
from app.core.excel import TIPO_XLSX
from app.models import Auditoria
from app.reports.definiciones import REPORTES
from app.reports.definiciones._comunes import porcentaje
from app.reports.motor import DefinicionReporte, normalizar_filtros
from app.schemas.categorias import CategoriaCrear
from app.schemas.productos import ProductoCrear
from app.schemas.proveedores import ProveedorCrear
from app.schemas.reportes import FiltrosReporte, TipoColumna
from app.services import categorias as servicio_categorias
from app.services import productos as servicio_productos
from app.services import proveedores as servicio_proveedores

API = "/api/v1/reportes"


@pytest.fixture
def admin(crear_usuario):
    return crear_usuario("admin")


@pytest.fixture
def h(headers_de, admin):
    return headers_de(admin)


@pytest.fixture
def stock(db, admin):
    """Árbol R > S > T, R > U y W; un proveedor en pesos y otro en dólares
    (dólar 1000). Solo cuentan los productos en T (3 u.), U (2 u.) y W (5 u.):
    el resto cae por inactivo, sin stock o con stock infinito."""
    def categoria(nombre, padre=None):
        return servicio_categorias.crear_categoria(
            db, admin, CategoriaCrear(nombre=nombre, parent_id=padre.id if padre else None), None)

    r = categoria("Rep R")
    s = categoria("Rep S", r)
    t = categoria("Rep T", s)
    u = categoria("Rep U", r)
    w = categoria("Rep W")
    ars = servicio_proveedores.crear_proveedor(db, admin, ProveedorCrear(nombre="Rep pesos", moneda="ars"), None)
    usd = servicio_proveedores.crear_proveedor(
        db, admin, ProveedorCrear(nombre="Rep dolares", moneda="usd", dolar_actual="1000"), None)

    def alta(descripcion, cat, proveedor, costo, stock_inicial=0, **extra):
        return servicio_productos.crear_producto(db, admin, ProductoCrear(
            descripcion=descripcion, categoria_id=cat.id, proveedor_id=proveedor.id, costo=costo, margen="50",
            stock_inicial=stock_inicial, **extra), None)

    productos = {"t": alta("Rep en T", t, ars, "1000", 3), "u": alta("Rep en U", u, usd, "2.50", 2),
                 "w": alta("Rep en W", w, ars, "400", 5)}
    alta("Rep inactivo", w, ars, "100", 4).activo = False
    alta("Rep sin stock", w, ars, "100")
    # Solo test: un producto de stock infinito con stock cargado no debe valorizarse.
    alta("Rep infinito", w, ars, "100", stock_infinito=True).stock = 7
    db.flush()
    return {"categorias": {"r": r, "s": s, "t": t, "u": u, "w": w}, "ars": ars, "usd": usd, "productos": productos}


def _pedir(client, h, clave, **params):
    r = client.get(f"{API}/{clave}", params=params, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _exportaciones(db) -> int:
    return db.scalar(select(func.count()).select_from(Auditoria).where(Auditoria.accion == "reporte.exportar"))


# --- Catálogo, permisos y validación ------------------------------------------

def test_catalogo_lista_los_reportes_con_sus_filtros(client, h):
    r = client.get(API, headers=h)
    assert r.status_code == 200
    stock = next(x for x in r.json() if x["clave"] == "stock-valorizado")
    assert stock["agrupaciones"] == ["categoria", "proveedor"]
    assert "desde" not in stock["filtros_admitidos"] and "proveedor_id" in stock["filtros_admitidos"]


@pytest.mark.parametrize("ruta", ["", "/opciones", "/stock-valorizado", "/stock-valorizado/excel",
                                  "/stock-valorizado/pdf"])
def test_vendedor_sin_reportes_ver_403(client, headers_de, crear_usuario, ruta):
    assert client.get(f"{API}{ruta}", headers=headers_de(crear_usuario("vendedor"))).status_code == 403


def test_supervisor_ve_reportes(client, headers_de, crear_usuario):
    assert client.get(f"{API}/stock-valorizado", headers=headers_de(crear_usuario("supervisor"))).status_code == 200


@pytest.mark.parametrize("sufijo", ["", "/excel", "/pdf"])
def test_clave_inexistente_404(client, h, sufijo):
    assert client.get(f"{API}/no-existe{sufijo}", headers=h).status_code == 404


@pytest.mark.parametrize("params", [
    {"desde": "2026-01-01"},               # stock valorizado no usa fechas
    {"usuario_id": 1},                     # ni vendedor
    {"agrupar": "dia"},                    # agrupación que no admite
    {"agrupar": "proveedor", "nivel": 2},  # nivel sin agrupar por categoría
    {"nivel": 6},                          # fuera de 1-5
    {"proveedor_id": 2**63},               # mayor que BIGINT
    {"inventado": "x"},                    # parámetro desconocido
])
def test_filtro_no_admitido_422(client, h, params):
    assert client.get(f"{API}/stock-valorizado", params=params, headers=h).status_code == 422


_CON_FECHAS = DefinicionReporte(clave="prueba", nombre="Prueba", descripcion="", filtros_admitidos=("desde", "hasta"),
                                columnas=lambda f: [], consulta=lambda db, f: ([], {}))


def test_fechas_por_defecto_del_primero_del_mes_a_hoy():
    f = normalizar_filtros(_CON_FECHAS, FiltrosReporte(), hoy=date(2026, 9, 29))
    assert (f.desde, f.hasta) == (date(2026, 9, 1), date(2026, 9, 29))


def test_fechas_por_defecto_desde_el_primero_del_mes_de_hasta():
    # Si solo llega "hasta" de un mes anterior, "desde" acompaña ese mes en
    # vez del mes actual (si no, quedaría posterior a "hasta" y daría 422).
    f = normalizar_filtros(_CON_FECHAS, FiltrosReporte(hasta=date(2026, 8, 15)), hoy=date(2026, 9, 29))
    assert (f.desde, f.hasta) == (date(2026, 8, 1), date(2026, 8, 15))


def test_rango_de_366_dias_se_acepta_y_de_367_da_409():
    normalizar_filtros(_CON_FECHAS, FiltrosReporte(desde=date(2024, 1, 1), hasta=date(2024, 12, 31)))  # bisiesto
    with pytest.raises(ReglaDeNegocio, match="366 días"):
        normalizar_filtros(_CON_FECHAS, FiltrosReporte(desde=date(2024, 1, 1), hasta=date(2025, 1, 1)))


def test_desde_posterior_a_hasta_422():
    with pytest.raises(NoProcesable):
        normalizar_filtros(_CON_FECHAS, FiltrosReporte(desde=date(2026, 9, 2), hasta=date(2026, 9, 1)))


# --- Gráfico --------------------------------------------------------------

@pytest.mark.parametrize("clave", list(REPORTES))
def test_cada_reporte_declara_un_grafico_sobre_sus_columnas(clave):
    definicion = REPORTES[clave]
    for agrupar in definicion.agrupaciones or (None,):
        filtros = normalizar_filtros(definicion, FiltrosReporte(agrupar=agrupar), hoy=date(2026, 9, 29))
        tipos = {c.clave: c.tipo for c in definicion.columnas(filtros)}
        grafico = definicion.grafico(filtros)
        assert grafico.tipo in ("linea", "barras")
        assert grafico.etiquetas and set(grafico.etiquetas) <= set(tipos)
        assert tipos[grafico.valor] in (TipoColumna.ENTERO, TipoColumna.MONEDA, TipoColumna.PORCENTAJE)
        assert grafico.titulo


# --- Stock valorizado ---------------------------------------------------------

def test_por_defecto_agrupa_por_categoria_de_nivel_1(client, h, stock):
    datos = _pedir(client, h, "stock-valorizado")
    assert datos["filtros"]["agrupar"] == "categoria" and datos["filtros"]["nivel"] == 1
    assert datos["columnas"][0] == {"clave": "grupo", "titulo": "Categoría", "tipo": "texto"}
    assert {f["grupo"]: f["unidades"] for f in datos["filas"]} == {"Rep R": 5, "Rep W": 5}


def test_por_proveedor_valoriza_a_costo_en_pesos_y_a_venta(client, h, stock):
    t, u, w = (stock["productos"][k] for k in "tuw")
    filas = {f["grupo"]: f for f in _pedir(client, h, "stock-valorizado", agrupar="proveedor")["filas"]}
    assert set(filas) == {"Rep pesos", "Rep dolares"}
    pesos = filas["Rep pesos"]
    assert pesos["productos"] == 2 and pesos["unidades"] == 8
    assert Decimal(pesos["valorizado_costo"]) == t.costo_en_pesos * 3 + w.costo_en_pesos * 5
    assert Decimal(pesos["valorizado_venta"]) == t.precio_venta * 3 + w.precio_venta * 5
    dolares = filas["Rep dolares"]
    assert dolares["valorizado_costo"] == "5000.00"                  # 2,50 × 1000 × 2
    assert Decimal(dolares["valorizado_venta"]) == u.precio_venta * 2
    assert Decimal(dolares["ganancia_potencial"]) == Decimal(dolares["valorizado_venta"]) - Decimal("5000.00")


def test_por_categoria_lleva_cada_producto_a_su_ancestro_del_nivel(client, h, stock):
    def por_nivel(nivel):
        filas = _pedir(client, h, "stock-valorizado", agrupar="categoria", nivel=nivel)["filas"]
        return {f["grupo"]: f["unidades"] for f in filas}

    assert por_nivel(2) == {"Rep R > Rep S": 3, "Rep R > Rep U": 2, "Rep W": 5}
    # Las ramas más cortas que el nivel pedido quedan en su propia categoría.
    assert por_nivel(3) == {"Rep R > Rep S > Rep T": 3, "Rep R > Rep U": 2, "Rep W": 5}


def test_filtra_por_rama_y_por_proveedor(client, h, stock):
    rama = _pedir(client, h, "stock-valorizado", agrupar="proveedor", categoria_id=stock["categorias"]["r"].id)
    assert {f["grupo"]: f["unidades"] for f in rama["filas"]} == {"Rep pesos": 3, "Rep dolares": 2}
    proveedor = _pedir(client, h, "stock-valorizado", agrupar="proveedor", proveedor_id=stock["usd"].id)
    assert [f["grupo"] for f in proveedor["filas"]] == ["Rep dolares"]


def test_totales_calculados_sobre_totales(client, h, stock):
    datos = _pedir(client, h, "stock-valorizado", agrupar="proveedor")
    totales = datos["totales"]
    assert totales["grupo"] is None and totales["productos"] == 3 and totales["unidades"] == 10
    assert Decimal(totales["valorizado_costo"]) == sum(Decimal(f["valorizado_costo"]) for f in datos["filas"])
    ganancia = Decimal(totales["valorizado_venta"]) - Decimal(totales["valorizado_costo"])
    assert Decimal(totales["margen_potencial"]) == porcentaje(ganancia, Decimal(totales["valorizado_venta"]))


# --- Exportaciones y auditoría ------------------------------------------------

def test_consultar_en_pantalla_no_audita(client, h, db, stock):
    _pedir(client, h, "stock-valorizado")
    assert _exportaciones(db) == 0


def test_excel_con_los_mismos_totales_que_el_json_y_auditado(client, h, db, stock, admin):
    datos = _pedir(client, h, "stock-valorizado", agrupar="proveedor")
    r = client.get(f"{API}/stock-valorizado/excel", params={"agrupar": "proveedor"}, headers=h)
    assert r.status_code == 200 and r.headers["content-type"].startswith(TIPO_XLSX)
    assert 'attachment; filename="reporte-stock-valorizado-' in r.headers["content-disposition"]
    hoja = load_workbook(BytesIO(r.content)).active
    ultima = [c.value for c in hoja[hoja.max_row]]
    assert ultima[0] == "Total"
    for indice, columna in enumerate(datos["columnas"][1:], start=1):
        assert Decimal(str(ultima[indice])) == Decimal(str(datos["totales"][columna["clave"]]))
    registro = db.scalars(select(Auditoria).where(Auditoria.accion == "reporte.exportar")).one()
    assert registro.usuario_id == admin.id and registro.entidad == "reporte" and registro.entidad_id is None
    assert registro.estado_nuevo == {"clave": "stock-valorizado", "formato": "excel", "filtros": {"agrupar": "proveedor"}}


def test_pdf_inline_y_auditado(client, h, db, stock):
    r = client.get(f"{API}/stock-valorizado/pdf", headers=h)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert r.headers["content-disposition"].startswith('inline; filename="reporte-stock-valorizado-')
    registro = db.scalars(select(Auditoria).where(Auditoria.accion == "reporte.exportar")).one()
    assert registro.estado_nuevo["formato"] == "pdf"
    assert registro.estado_nuevo["filtros"] == {"agrupar": "categoria", "nivel": 1}


def test_pdf_sin_datos_se_genera_igual(client, h):
    r = client.get(f"{API}/stock-valorizado/pdf", params={"proveedor_id": 2**62}, headers=h)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


# --- Opciones -----------------------------------------------------------------

def test_opciones_devuelve_los_catalogos_de_los_filtros(client, h, stock):
    r = client.get(f"{API}/opciones", headers=h)
    assert r.status_code == 200
    datos = r.json()
    assert set(datos) == {"proveedores", "categorias", "vendedores", "medios_de_pago"}
    assert {"id": stock["usd"].id, "nombre": "Rep dolares"} in datos["proveedores"]
    assert {"id": stock["categorias"]["t"].id, "nombre": "Rep R > Rep S > Rep T", "nivel": 3} in datos["categorias"]
