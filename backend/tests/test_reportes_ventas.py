"""Reportes de ventas (sesión 09, Task 2). Las ventas se arman a mano (sin
services.ventas) para controlar fechas e importes y verificar cada número.
Ver la tabla de datos en el plan del módulo 09 (Task 2)."""
import itertools
from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import EstadoVenta, MedioDePago, PlanCuotas, Venta, VentaItem, VentaPago
from app.schemas.categorias import CategoriaCrear
from app.schemas.productos import ProductoCrear
from app.schemas.proveedores import ProveedorCrear
from app.services import categorias as servicio_categorias
from app.services import productos as servicio_productos
from app.services import proveedores as servicio_proveedores

API = "/api/v1/reportes"
RANGO = {"desde": "2025-03-01", "hasta": "2025-04-30"}
D = Decimal
_numeros = itertools.count(1)


def _venta(db, usuario, items, pagos, *, confirmada_en=None, anulada=None, en_curso=False):
    """`items` = [(producto, precio_lista, precio_final, costo_unitario)];
    `pagos` = [(medio, plan o None, monto)] (el recargo sale del plan);
    `anulada` = (fecha de anulación, usuario que anula)."""
    n = next(_numeros)
    cobros = []
    for medio, plan, monto in pagos:
        recargo = (D(monto) * plan.recargo_cliente / 100).quantize(D("0.01")) if plan else D("0")
        cobros.append((medio, plan, D(monto), recargo))
    subtotal = sum((D(i[1]) for i in items), D("0"))
    cobrable = sum((D(i[2]) for i in items), D("0"))
    recargo_total = sum((c[3] for c in cobros), D("0"))
    venta = Venta(numero=f"RT-{n:06d}", usuario_id=usuario.id, subtotal=subtotal, descuento_total=subtotal - cobrable,
                  recargo_total=recargo_total, total=cobrable + recargo_total,
                  estado=EstadoVenta.EN_CURSO if en_curso else EstadoVenta.CONFIRMADA)
    if not en_curso:
        venta.codigo_cambio = f"RT{n:06d}"
        venta.confirmada_en = confirmada_en
    if anulada:
        venta.estado = EstadoVenta.ANULADA
        venta.anulada_en, quien = anulada
        venta.anulada_por_id = quien.id
        venta.motivo_anulacion = "Prueba de reportes"
    db.add(venta)
    db.flush()
    for orden, (producto, lista, final, costo) in enumerate(items):
        db.add(VentaItem(venta_id=venta.id, producto_id=producto.id, precio_lista=D(lista), precio_unitario=D(lista),
                         precio_final=D(final), costo_unitario=D(costo), en_promocion=D(final) < D(lista), orden=orden))
    for medio, plan, monto, recargo in cobros:
        db.add(VentaPago(venta_id=venta.id, medio_de_pago_id=medio.id, plan_cuotas_id=plan.id if plan else None,
                         monto=monto, recargo=recargo, monto_total=monto + recargo))
    db.flush()
    return venta


@pytest.fixture
def e(db, crear_usuario):
    admin = crear_usuario("admin")
    ana = crear_usuario("vendedor", nombre="Ana Rep")
    beto = crear_usuario("vendedor", nombre="Beto Rep")

    def categoria(nombre, padre=None):
        return servicio_categorias.crear_categoria(
            db, admin, CategoriaCrear(nombre=nombre, parent_id=padre.id if padre else None), None)

    r = categoria("Rep R")
    s = categoria("Rep S", r)
    t = categoria("Rep T", s)
    u = categoria("Rep U", r)
    w = categoria("Rep W")
    ars = servicio_proveedores.crear_proveedor(db, admin, ProveedorCrear(nombre="Rep pesos", moneda="ars"), None)
    usd = servicio_proveedores.crear_proveedor(
        db, admin, ProveedorCrear(nombre="Rep dolares", moneda="usd", dolar_actual="1000"), None)

    def alta(descripcion, cat, proveedor, costo):
        return servicio_productos.crear_producto(db, admin, ProductoCrear(
            descripcion=descripcion, categoria_id=cat.id, proveedor_id=proveedor.id, costo=costo, margen="0"), None)

    pa, pb, pd = alta("Rep A", t, ars, "1000"), alta("Rep B", w, ars, "500"), alta("Rep D", u, usd, "10")
    contado = MedioDePago(nombre="Contado Rep")
    tarjeta = MedioDePago(nombre="Tarjeta Rep", soporta_cuotas=True)
    db.add_all([contado, tarjeta])
    db.flush()
    plan = PlanCuotas(medio_de_pago_id=tarjeta.id, cuotas=3, recargo_cliente=D("10"), costo_medio=D("5"))
    db.add(plan)
    db.flush()

    a = (pa, "2000", "2000", "1000")
    _venta(db, ana, [a, (pa, "2000", "1800", "1000"), a, (pd, "20000", "20000", "10000")],
           [(contado, None, "25800")], confirmada_en=datetime(2025, 3, 3, 10))
    _venta(db, beto, [(pb, "1000", "1000", "500")], [(tarjeta, plan, "1000")], confirmada_en=datetime(2025, 3, 3, 18))
    _venta(db, ana, [(pd, "20000", "19000", "10000")], [(contado, None, "10000"), (tarjeta, None, "9000")],
           confirmada_en=datetime(2025, 3, 20, 12))
    _venta(db, beto, [a], [(contado, None, "2000")], confirmada_en=datetime(2025, 3, 5, 12),
           anulada=(datetime(2025, 4, 2, 11), admin))
    _venta(db, ana, [(pb, "1000", "1000", "500")], [], en_curso=True)
    _venta(db, ana, [(pb, "1000", "1000", "500")], [(contado, None, "1000")], confirmada_en=datetime(2025, 2, 27, 12))
    return {"admin": admin, "ana": ana, "beto": beto, "cat": {"r": r, "w": w}, "ars": ars, "usd": usd,
            "pa": pa, "pb": pb, "pd": pd, "contado": contado, "tarjeta": tarjeta}


@pytest.fixture
def h(headers_de, e):
    return headers_de(e["admin"])


def _pedir(client, h, clave, **params):
    r = client.get(f"{API}/{clave}", params={**RANGO, **params}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _sin(fila, *claves):
    return {k: v for k, v in fila.items() if k not in claves}


# --- Ventas por período -------------------------------------------------------

def test_ventas_por_dia_con_anuladas_por_fecha_de_anulacion(client, h, e):
    datos = _pedir(client, h, "ventas-periodo")
    assert [c["tipo"] for c in datos["columnas"][:1]] == ["fecha"]
    assert datos["filas"] == [
        {"periodo": "2025-03-03", "ventas": 2, "unidades": 5, "subtotal": "27000.00", "descuentos": "200.00",
         "recargos": "100.00", "total": "26900.00", "ticket_promedio": "13450.00", "anuladas": 0,
         "anuladas_monto": "0.00"},
        {"periodo": "2025-03-20", "ventas": 1, "unidades": 1, "subtotal": "20000.00", "descuentos": "1000.00",
         "recargos": "0.00", "total": "19000.00", "ticket_promedio": "19000.00", "anuladas": 0,
         "anuladas_monto": "0.00"},
        {"periodo": "2025-04-02", "ventas": 0, "unidades": 0, "subtotal": "0.00", "descuentos": "0.00",
         "recargos": "0.00", "total": "0.00", "ticket_promedio": "0.00", "anuladas": 1, "anuladas_monto": "2000.00"},
    ]
    assert datos["totales"] == {
        "periodo": None, "ventas": 3, "unidades": 6, "subtotal": "47000.00", "descuentos": "1200.00",
        "recargos": "100.00", "total": "45900.00", "ticket_promedio": "15300.00", "anuladas": 1,
        "anuladas_monto": "2000.00"}


def test_ventas_por_mes(client, h, e):
    datos = _pedir(client, h, "ventas-periodo", agrupar="mes")
    assert datos["columnas"][0]["tipo"] == "mes"
    assert [(f["periodo"], f["ventas"], f["total"], f["anuladas"]) for f in datos["filas"]] == [
        ("2025-03-01", 3, "45900.00", 0), ("2025-04-01", 0, "0.00", 1)]


def test_ventas_periodo_no_admite_proveedor_ni_categoria(client, h, e):
    for params in ({"proveedor_id": e["ars"].id}, {"categoria_id": e["cat"]["r"].id}):
        assert client.get(f"{API}/ventas-periodo", params={**RANGO, **params}, headers=h).status_code == 422


def test_rango_de_mas_de_366_dias_409(client, h, e):
    r = client.get(f"{API}/ventas-periodo", params={"desde": "2024-01-01", "hasta": "2025-01-01"}, headers=h)
    assert r.status_code == 409


# --- Medios de pago -----------------------------------------------------------

def test_medios_de_pago_con_plan_recargo_y_costo_del_medio(client, h, e):
    datos = _pedir(client, h, "medios-pago")
    assert datos["filas"] == [
        {"medio": "Contado Rep", "plan": "Un pago", "pagos": 2, "monto": "35800.00", "recargo": "0.00",
         "total_cobrado": "35800.00", "costo_medio": "0.00", "neto": "35800.00"},
        {"medio": "Tarjeta Rep", "plan": "Un pago", "pagos": 1, "monto": "9000.00", "recargo": "0.00",
         "total_cobrado": "9000.00", "costo_medio": "0.00", "neto": "9000.00"},
        {"medio": "Tarjeta Rep", "plan": "3 cuotas +10%", "pagos": 1, "monto": "1000.00", "recargo": "100.00",
         "total_cobrado": "1100.00", "costo_medio": "55.00", "neto": "1045.00"},   # 1100 × 5 %
    ]
    assert _sin(datos["totales"], "medio", "plan") == {
        "pagos": 4, "monto": "45800.00", "recargo": "100.00", "total_cobrado": "45900.00", "costo_medio": "55.00",
        "neto": "45845.00"}


def test_medios_de_pago_filtra_por_vendedor_y_por_medio(client, h, e):
    beto = _pedir(client, h, "medios-pago", usuario_id=e["beto"].id)["filas"]
    assert [(f["medio"], f["plan"]) for f in beto] == [("Tarjeta Rep", "3 cuotas +10%")]
    contado = _pedir(client, h, "medios-pago", medio_de_pago_id=e["contado"].id)["filas"]
    assert [f["medio"] for f in contado] == ["Contado Rep"]


# --- Vendedores ---------------------------------------------------------------

def test_vendedores(client, h, e):
    datos = _pedir(client, h, "vendedores")
    assert datos["filas"] == [
        {"vendedor": "Ana Rep", "ventas": 2, "unidades": 5, "total": "44800.00", "ticket_promedio": "22400.00",
         "descuentos": "1200.00"},
        {"vendedor": "Beto Rep", "ventas": 1, "unidades": 1, "total": "1100.00", "ticket_promedio": "1100.00",
         "descuentos": "0.00"},
    ]
    assert datos["totales"] == {"vendedor": None, "ventas": 3, "unidades": 6, "total": "45900.00",
                                "ticket_promedio": "15300.00", "descuentos": "1200.00"}


def test_vendedores_no_admite_medio(client, h, e):
    r = client.get(f"{API}/vendedores", params={**RANGO, "medio_de_pago_id": e["contado"].id}, headers=h)
    assert r.status_code == 422


# --- Productos más vendidos ---------------------------------------------------

def test_productos_por_unidades_y_por_importe(client, h, e):
    pa, pb, pd = e["pa"], e["pb"], e["pd"]
    por_unidades = _pedir(client, h, "productos-vendidos")
    assert [c["clave"] for c in por_unidades["columnas"]] == ["codigo", "descripcion", "unidades", "importe", "porcentaje"]
    assert por_unidades["filas"] == [
        {"codigo": pa.codigo_con_verificador, "descripcion": "Rep A", "unidades": 3, "importe": "5800.00",
         "porcentaje": "12.66"},
        {"codigo": pd.codigo_con_verificador, "descripcion": "Rep D", "unidades": 2, "importe": "39000.00",
         "porcentaje": "85.15"},
        {"codigo": pb.codigo_con_verificador, "descripcion": "Rep B", "unidades": 1, "importe": "1000.00",
         "porcentaje": "2.18"},
    ]
    assert por_unidades["totales"] == {"codigo": None, "descripcion": None, "unidades": 6, "importe": "45800.00",
                                       "porcentaje": "100.00"}
    por_importe = _pedir(client, h, "productos-vendidos", orden="importe")
    assert [f["descripcion"] for f in por_importe["filas"]] == ["Rep D", "Rep A", "Rep B"]


def test_productos_limite_no_achica_los_totales(client, h, e):
    datos = _pedir(client, h, "productos-vendidos", limite=1)
    assert [f["descripcion"] for f in datos["filas"]] == ["Rep A"]
    assert datos["totales"]["importe"] == "45800.00" and datos["totales"]["unidades"] == 6
    assert client.get(f"{API}/productos-vendidos", params={**RANGO, "limite": 501}, headers=h).status_code == 422


def test_productos_por_categoria_de_nivel(client, h, e):
    nivel_1 = _pedir(client, h, "productos-vendidos", agrupar="categoria")
    assert [c["clave"] for c in nivel_1["columnas"]] == ["categoria", "unidades", "importe", "porcentaje"]
    assert nivel_1["filas"] == [
        {"categoria": "Rep R", "unidades": 5, "importe": "44800.00", "porcentaje": "97.82"},
        {"categoria": "Rep W", "unidades": 1, "importe": "1000.00", "porcentaje": "2.18"},
    ]
    nivel_2 = _pedir(client, h, "productos-vendidos", agrupar="categoria", nivel=2)
    assert [(f["categoria"], f["unidades"]) for f in nivel_2["filas"]] == [
        ("Rep R > Rep S", 3), ("Rep R > Rep U", 2), ("Rep W", 1)]


def test_productos_filtra_por_proveedor_y_por_rama(client, h, e):
    pesos = _pedir(client, h, "productos-vendidos", proveedor_id=e["ars"].id)
    assert [(f["descripcion"], f["porcentaje"]) for f in pesos["filas"]] == [("Rep A", "85.29"), ("Rep B", "14.71")]
    rama = _pedir(client, h, "productos-vendidos", categoria_id=e["cat"]["r"].id)
    assert [f["descripcion"] for f in rama["filas"]] == ["Rep A", "Rep D"]


# --- Rentabilidad -------------------------------------------------------------

def test_rentabilidad_por_proveedor_sin_recargos(client, h, e):
    datos = _pedir(client, h, "rentabilidad", agrupar="proveedor")
    assert datos["filas"] == [
        {"grupo": "Rep dolares", "unidades": 2, "importe": "39000.00", "costo": "20000.00", "ganancia": "19000.00",
         "margen": "48.72"},
        {"grupo": "Rep pesos", "unidades": 4, "importe": "6800.00", "costo": "3500.00", "ganancia": "3300.00",
         "margen": "48.53"},
    ]
    # Importe = Σ precio_final (45800): el recargo de 100 de V2 no entra.
    assert datos["totales"] == {"grupo": None, "unidades": 6, "importe": "45800.00", "costo": "23500.00",
                                "ganancia": "22300.00", "margen": "48.69"}


def test_rentabilidad_por_dia_mes_y_categoria(client, h, e):
    por_dia = _pedir(client, h, "rentabilidad")
    assert datos_de(por_dia) == [("2025-03-03", 5, "26800.00", "13500.00", "49.63"),
                                 ("2025-03-20", 1, "19000.00", "10000.00", "47.37")]
    assert datos_de(_pedir(client, h, "rentabilidad", agrupar="mes")) == [
        ("2025-03-01", 6, "45800.00", "23500.00", "48.69")]
    assert datos_de(_pedir(client, h, "rentabilidad", agrupar="categoria")) == [
        ("Rep R", 5, "44800.00", "23000.00", "48.66"), ("Rep W", 1, "1000.00", "500.00", "50.00")]


def datos_de(respuesta):
    return [(f["grupo"], f["unidades"], f["importe"], f["costo"], f["margen"]) for f in respuesta["filas"]]


def test_rentabilidad_usa_el_costo_congelado(client, h, e, db):
    antes = _pedir(client, h, "rentabilidad", agrupar="proveedor")["totales"]
    e["usd"].dolar_actual = D("2000")
    e["pa"].costo = D("5000")
    db.flush()
    assert _pedir(client, h, "rentabilidad", agrupar="proveedor")["totales"] == antes


def test_rentabilidad_no_admite_agrupar_por_producto(client, h, e):
    assert client.get(f"{API}/rentabilidad", params={**RANGO, "agrupar": "producto"}, headers=h).status_code == 422


# --- Exportaciones y opciones -------------------------------------------------

@pytest.mark.parametrize("clave", ["ventas-periodo", "medios-pago", "vendedores", "productos-vendidos",
                                   "rentabilidad", "stock-valorizado"])
def test_excel_y_pdf_de_cada_reporte(client, h, e, clave):
    params = {} if clave == "stock-valorizado" else RANGO
    datos = client.get(f"{API}/{clave}", params=params, headers=h).json()
    r = client.get(f"{API}/{clave}/excel", params=params, headers=h)
    assert r.status_code == 200
    hoja = load_workbook(BytesIO(r.content)).active
    ultima = [c.value for c in hoja[hoja.max_row]]
    assert ultima[0] == "Total"
    for indice, columna in enumerate(datos["columnas"][1:], start=1):
        valor = datos["totales"][columna["clave"]]
        if valor is not None:
            assert Decimal(str(ultima[indice])) == Decimal(str(valor))
    pdf = client.get(f"{API}/{clave}/pdf", params=params, headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_nombre_del_excel_lleva_el_rango(client, h, e):
    r = client.get(f"{API}/ventas-periodo/excel", params=RANGO, headers=h)
    assert 'filename="reporte-ventas-periodo-2025-03-01_2025-04-30.xlsx"' in r.headers["content-disposition"]


def test_catalogo_en_el_orden_del_hub(client, h, e):
    assert [x["clave"] for x in client.get(API, headers=h).json()] == [
        "ventas-periodo", "medios-pago", "vendedores", "productos-vendidos", "rentabilidad", "stock-valorizado"]


def test_opciones_vendedores_que_vendieron_aunque_esten_inactivos(client, h, e, db, crear_usuario):
    crear_usuario("vendedor", nombre="Sin ventas Rep")
    e["beto"].activo = False
    e["tarjeta"].activo = False
    db.flush()
    datos = client.get(f"{API}/opciones", headers=h).json()
    nombres = [v["nombre"] for v in datos["vendedores"]]
    assert "Ana Rep" in nombres and "Beto Rep" in nombres and "Sin ventas Rep" not in nombres
    assert "Tarjeta Rep" in [m["nombre"] for m in datos["medios_de_pago"]]
